import hashlib
import importlib.util
import sys
import unittest
from collections import Counter
from dataclasses import FrozenInstanceError, fields
from pathlib import Path

from openpyxl import load_workbook


PYTHON_ROOT = Path(__file__).resolve().parents[1]
REPOSITORY_ROOT = Path(__file__).resolve().parents[3]
MODULE_DIRECTORY = PYTHON_ROOT / "tools" / "canonical_export"
SCRIPT_PATH = MODULE_DIRECTORY / "canonical_validation_rules.py"
WORKBOOKS = (
    REPOSITORY_ROOT / "Aeterna dokumentációk" / "REGISTRY.xlsx",
    REPOSITORY_ROOT / "Aeterna dokumentációk" / "CARDDATABASE.xlsx",
)


def load_rules_module():
    module_name = "canonical_validation_rules"
    sys.path.insert(0, str(MODULE_DIRECTORY))
    try:
        spec = importlib.util.spec_from_file_location(module_name, SCRIPT_PATH)
        module = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = module
        spec.loader.exec_module(module)
        return module
    finally:
        sys.path.remove(str(MODULE_DIRECTORY))


rules = load_rules_module()


def make_record(kind="custom_expression", rule_id="test_rule", **overrides):
    record = {field: "#NULL" for field in rules.RULE_RECORD_FIELDS}
    record.update(
        {
            "validation_rule_id": rule_id,
            "rule_scope_id": "field",
            "target_table_id": "test_table",
            "target_field_id": "fld_test_table_value",
            "validation_kind_id": kind,
            "condition_expression": "value == value",
            "severity_id": "error",
            "blocking": True,
            "error_code": "TEST_RULE_FAILURE",
            "message": "Synthetic contract fixture.",
            "validation_stage_id": "pre_export",
            "status": "active",
            "source_id": "src_test",
            "source_ref": "test_table.value",
        }
    )
    record.update(overrides)
    return record


def positive_kind_record(kind, index=0):
    rule_id = f"positive_{index:02d}_{kind}"
    common = {"validation_rule_id": rule_id}
    if kind == "allowed_value":
        return make_record(
            kind,
            **common,
            condition_expression="#NULL",
            operator_id="op_in",
            comparison_value="test_group",
        )
    if kind == "contract_field_invariant":
        return make_record(kind, **common)
    if kind == "contract_instance_consistency":
        return make_record(
            kind,
            **common,
            rule_scope_id="contract_instance",
            target_field_id="#NULL",
            comparison_value="payload_contract_v1",
        )
    if kind == "cross_field_consistency":
        return make_record(
            kind,
            **common,
            rule_scope_id="cross_table",
            reference_table_id="reference_table",
            reference_field_id="fld_reference_table_id",
        )
    if kind in {"custom_expression", "definition_invariant"}:
        return make_record(kind, **common)
    if kind == "hierarchy_integrity":
        return make_record(
            kind,
            **common,
            condition_expression=(
                'no_self_reference("id","parent_id") and '
                'hierarchy_is_acyclic("id","parent_id")'
            ),
        )
    if kind == "normalized_uniqueness":
        return make_record(
            kind,
            **common,
            condition_expression="for_each_group(group_id, value == value)",
        )
    if kind == "qualified_reference_match":
        return make_record(kind, **common)
    if kind == "range":
        return make_record(
            kind,
            **common,
            condition_expression="#NULL",
            operator_id="op_greater_than_or_equal",
            comparison_value=1.0,
            minimum_value=1.0,
        )
    if kind == "reference_integrity":
        return make_record(
            kind,
            **common,
            condition_expression="#NULL",
            reference_table_id="reference_table",
            reference_field_id="fld_reference_table_id",
        )
    if kind == "target_group_membership":
        return make_record(
            kind,
            **common,
            reference_table_id="reference_table",
            reference_field_id="fld_reference_table_id",
        )
    if kind == "uniqueness":
        return make_record(
            kind,
            **common,
            rule_scope_id="table",
            condition_expression='unique_by(["value"])',
        )
    raise AssertionError(f"Missing positive fixture for {kind}.")


def diagnostic_codes(result):
    return tuple(item.code for item in result.diagnostics)


def read_active_rule_inputs():
    hashes_before = {
        path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
    }
    headers_by_workbook = {}
    inputs = []
    for path in WORKBOOKS:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = workbook["VALIDATION_RULES"].iter_rows(values_only=True)
            headers = tuple(next(rows))
            headers_by_workbook[path.name] = headers
            for row_number, row in enumerate(rows, start=2):
                record = dict(zip(headers, row))
                if record["status"] == "active":
                    inputs.append(
                        rules.ValidationRuleInput(record, path.name, row_number)
                    )
        finally:
            workbook.close()
    hashes_after = {
        path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
    }
    if hashes_before != hashes_after:
        raise AssertionError("Read-only catalog fixture changed a workbook.")
    return headers_by_workbook, tuple(inputs)


def walk_ast(node):
    yield node
    if isinstance(node, rules.ListLiteral):
        for item in node.items:
            yield from walk_ast(item)
    elif isinstance(node, rules.UnaryOperation):
        yield from walk_ast(node.operand)
    elif isinstance(node, rules.BinaryOperation):
        yield from walk_ast(node.left)
        yield from walk_ast(node.right)
    elif isinstance(node, rules.Call):
        for argument in node.arguments:
            yield from walk_ast(argument)
    elif isinstance(node, rules.MemberAccess):
        yield from walk_ast(node.target)
    elif isinstance(node, rules.Binding):
        yield from walk_ast(node.value)
    elif isinstance(node, rules.AstSequence):
        for statement in node.statements:
            yield from walk_ast(statement)


class TestInputAndTypedModel(unittest.TestCase):
    def test_exact_workbook_record_schema_is_explicit(self):
        self.assertEqual(
            rules.RULE_RECORD_FIELDS,
            (
                "validation_rule_id",
                "rule_scope_id",
                "target_table_id",
                "target_field_id",
                "validation_kind_id",
                "condition_expression",
                "operator_id",
                "comparison_value",
                "minimum_value",
                "maximum_value",
                "reference_table_id",
                "reference_field_id",
                "severity_id",
                "blocking",
                "error_code",
                "message_key",
                "message",
                "validation_stage_id",
                "status",
                "source_id",
                "source_ref",
                "notes",
            ),
        )

    def test_unknown_and_missing_columns_are_not_silently_ignored(self):
        unexpected = make_record(unexpected_column="value")
        missing = make_record()
        del missing["notes"]
        for record, detail_name in (
            (unexpected, "unexpected_fields"),
            (missing, "missing_fields"),
        ):
            with self.subTest(detail=detail_name):
                result = rules.build_validation_rule_catalog((record,))
                self.assertFalse(result.is_valid)
                shape = next(
                    item
                    for item in result.diagnostics
                    if item.code == rules.VALIDATION_RULE_SHAPE_INVALID
                    and item.field_name == "record"
                )
                self.assertIn(detail_name, shape.as_dict()["context"])

    def test_null_and_tbd_sentinels_remain_distinct(self):
        structured = make_record(
            "uniqueness",
            rule_id="sentinel_null",
            rule_scope_id="table",
            condition_expression="#NULL",
        )
        bare_tbd = make_record(
            rule_id="sentinel_bare_tbd",
            condition_expression="value == tbd",
        )
        quoted_tbd = make_record(
            rule_id="sentinel_quoted_tbd",
            condition_expression='value == "#TBD"',
        )
        result = rules.build_validation_rule_catalog(
            (structured, bare_tbd, quoted_tbd)
        )
        self.assertTrue(result.is_valid, result.diagnostics)
        self.assertIsNone(result.catalog.get("sentinel_null").condition_ast)
        self.assertIsNotNone(result.catalog.get("sentinel_bare_tbd").condition_ast)
        self.assertIsNotNone(result.catalog.get("sentinel_quoted_tbd").condition_ast)
        self.assertNotEqual(
            result.catalog.get("sentinel_bare_tbd").condition_ast_hash,
            result.catalog.get("sentinel_quoted_tbd").condition_ast_hash,
        )

        workbook_tbd = make_record(
            rule_id="sentinel_workbook_tbd",
            condition_expression="#TBD",
        )
        invalid = rules.build_validation_rule_catalog((workbook_tbd,))
        self.assertTrue(
            any(item.code.startswith("VALIDATION_EXPR_") for item in invalid.diagnostics)
        )

    def test_definition_and_catalog_are_immutable_and_execution_free(self):
        result = rules.build_validation_rule_catalog((make_record(),))
        self.assertTrue(result.is_valid, result.diagnostics)
        definition = result.catalog.rules[0]
        with self.assertRaises(FrozenInstanceError):
            definition.rule_id = "changed"
        with self.assertRaises(FrozenInstanceError):
            result.catalog.rules = ()
        self.assertIsInstance(result.catalog.rules, tuple)
        names = {field.name for field in fields(rules.ValidationRuleDefinition)}
        self.assertFalse(
            names
            & {
                "result",
                "execution_result",
                "diagnostic_count",
                "pass_count",
                "rule_source_identity",
            }
        )


class TestKindShapeContracts(unittest.TestCase):
    def test_all_thirteen_current_kinds_have_positive_shape_coverage(self):
        records = tuple(
            positive_kind_record(kind, index)
            for index, kind in enumerate(sorted(rules.KNOWN_VALIDATION_KINDS))
        )
        result = rules.build_validation_rule_catalog(records)
        self.assertTrue(result.is_valid, result.diagnostics)
        self.assertEqual(len(result.catalog), 13)
        self.assertEqual(
            {rule.validation_kind_id for rule in result.catalog.rules},
            rules.KNOWN_VALIDATION_KINDS,
        )
        self.assertEqual(
            sum(len(contract.variants) for contract in rules.RULE_KIND_CONTRACTS),
            19,
        )

    def test_unknown_contract_tokens_are_structured(self):
        cases = (
            (
                "validation_kind_id",
                "future_kind",
                rules.VALIDATION_RULE_KIND_UNKNOWN,
            ),
            (
                "validation_stage_id",
                "future_stage",
                rules.VALIDATION_RULE_STAGE_UNKNOWN,
            ),
            (
                "severity_id",
                "fatal-ish",
                rules.VALIDATION_RULE_SEVERITY_UNKNOWN,
            ),
            ("status", "inactive-ish", rules.VALIDATION_RULE_STATUS_INVALID),
            ("blocking", "yes", rules.VALIDATION_RULE_BLOCKING_INVALID),
            ("blocking", 1, rules.VALIDATION_RULE_BLOCKING_INVALID),
        )
        for index, (field_name, value, expected) in enumerate(cases):
            with self.subTest(field=field_name, value=value):
                record = make_record(rule_id=f"invalid_token_{index}")
                record[field_name] = value
                result = rules.build_validation_rule_catalog((record,))
                self.assertIn(expected, diagnostic_codes(result))
                self.assertIsNone(result.catalog)

    def test_duplicate_rule_id_blocks_the_entire_catalog(self):
        first = make_record(rule_id="duplicate")
        second = make_record(rule_id="duplicate")
        result = rules.build_validation_rule_catalog((first, second))
        self.assertFalse(result.is_valid)
        self.assertIsNone(result.catalog)
        diagnostic = next(
            item
            for item in result.diagnostics
            if item.code == rules.VALIDATION_RULE_DUPLICATE_ID
        )
        self.assertEqual(diagnostic.as_dict()["context"]["duplicate_count"], 2)

    def test_required_and_forbidden_shape_fields_are_enforced(self):
        cases = (
            make_record(
                "custom_expression",
                rule_id="missing_expression",
                condition_expression="#NULL",
            ),
            make_record(
                "allowed_value",
                rule_id="unexpected_expression",
                operator_id="op_in",
                comparison_value="group",
            ),
            make_record(
                "uniqueness",
                rule_id="forbidden_maximum",
                rule_scope_id="table",
                condition_expression="#NULL",
                maximum_value=10.0,
            ),
        )
        for record in cases:
            with self.subTest(rule=record["validation_rule_id"]):
                result = rules.build_validation_rule_catalog((record,))
                self.assertIn(
                    rules.VALIDATION_RULE_SHAPE_INVALID,
                    diagnostic_codes(result),
                )

    def test_half_reference_pair_is_rejected(self):
        record = positive_kind_record("reference_integrity")
        record["reference_field_id"] = "#NULL"
        result = rules.build_validation_rule_catalog((record,))
        diagnostic = next(
            item
            for item in result.diagnostics
            if item.code == rules.VALIDATION_RULE_SHAPE_INVALID
            and item.field_name == "reference_table_id,reference_field_id"
        )
        self.assertIn("both be present", diagnostic.message)

    def test_uniqueness_expression_family_is_allowlisted(self):
        accepted = positive_kind_record("uniqueness")
        rejected = positive_kind_record("uniqueness")
        rejected["validation_rule_id"] = "invalid_uniqueness_family"
        rejected["condition_expression"] = "value == value"
        self.assertTrue(rules.build_validation_rule_catalog((accepted,)).is_valid)
        result = rules.build_validation_rule_catalog((rejected,))
        self.assertIn(rules.VALIDATION_RULE_SHAPE_INVALID, diagnostic_codes(result))

    def test_range_values_are_shape_checked_without_numeric_canonicalization(self):
        numeric = positive_kind_record("range")
        result = rules.build_validation_rule_catalog((numeric,))
        self.assertTrue(result.is_valid, result.diagnostics)
        definition = result.catalog.rules[0]
        self.assertIs(type(definition.comparison_value), float)
        self.assertIs(type(definition.minimum_value), float)

        invalid = positive_kind_record("range")
        invalid["comparison_value"] = "1.0"
        result = rules.build_validation_rule_catalog((invalid,))
        self.assertIn(rules.VALIDATION_RULE_SHAPE_INVALID, diagnostic_codes(result))


class TestExpressionAndSymbols(unittest.TestCase):
    def test_parser_diagnostic_code_and_detail_are_preserved_with_rule_context(self):
        record = make_record(
            rule_id="invalid_parser_expression",
            condition_expression="when(",
        )
        wrapped = rules.ValidationRuleInput(record, "TEST.xlsx", 27)
        result = rules.build_validation_rule_catalog((wrapped,))
        diagnostic = next(
            item for item in result.diagnostics if item.code.startswith("VALIDATION_EXPR_")
        )
        self.assertEqual(diagnostic.rule_id, "invalid_parser_expression")
        self.assertEqual(diagnostic.component_identity, "TEST.xlsx")
        self.assertEqual(diagnostic.source_row, 27)
        self.assertEqual(diagnostic.field_name, "condition_expression")
        self.assertIsNotNone(diagnostic.expression_diagnostic)
        self.assertEqual(
            diagnostic.expression_diagnostic.as_dict(),
            diagnostic.as_dict()["expression_diagnostic"],
        )

    def test_sequential_binding_is_visible_and_inventory_is_typed(self):
        record = make_record(
            rule_id="sequential_binding",
            condition_expression=(
                'row = lookup_record("items","id",item_id); row != null'
            ),
        )
        result = rules.build_validation_rule_catalog((record,))
        self.assertTrue(result.is_valid, result.diagnostics)
        definition = result.catalog.rules[0]
        self.assertEqual(definition.local_bindings, ("row",))
        self.assertEqual(definition.free_identifiers, ("item_id",))

    def test_forward_reference_is_rejected(self):
        record = make_record(
            rule_id="forward_binding",
            condition_expression=(
                'row != null; row = lookup_record("items","id",item_id)'
            ),
        )
        result = rules.build_validation_rule_catalog((record,))
        self.assertIn(
            rules.VALIDATION_RULE_BINDING_FORWARD_REFERENCE,
            diagnostic_codes(result),
        )

    def test_same_scope_rebinding_is_rejected(self):
        record = make_record(
            rule_id="duplicate_binding",
            condition_expression=(
                'row = lookup_record("items","id",first_id); '
                'row = lookup_record("items","id",second_id); row != null'
            ),
        )
        result = rules.build_validation_rule_catalog((record,))
        self.assertIn(
            rules.VALIDATION_RULE_BINDING_REBIND,
            diagnostic_codes(result),
        )

    def test_builtin_and_member_names_are_not_free_identifiers(self):
        record = make_record(
            rule_id="builtin_symbol_boundary",
            condition_expression=(
                'row = lookup_record("items","id",item_id); '
                'row.member_field == status'
            ),
        )
        result = rules.build_validation_rule_catalog((record,))
        self.assertTrue(result.is_valid, result.diagnostics)
        definition = result.catalog.rules[0]
        self.assertEqual(definition.free_identifiers, ("item_id", "status"))
        self.assertNotIn("lookup_record", definition.free_identifiers)
        self.assertNotIn("member_field", definition.free_identifiers)


class TestCatalogDeterminism(unittest.TestCase):
    def test_catalog_order_lookup_and_inventory_are_deterministic(self):
        records = tuple(
            positive_kind_record(kind, index)
            for index, kind in enumerate(sorted(rules.KNOWN_VALIDATION_KINDS))
        )
        forward = rules.build_validation_rule_catalog(records)
        reverse = rules.build_validation_rule_catalog(tuple(reversed(records)))
        self.assertTrue(forward.is_valid, forward.diagnostics)
        self.assertEqual(forward, reverse)
        expected_ids = tuple(
            sorted(
                (record["validation_rule_id"] for record in records),
                key=lambda value: value.encode("utf-8"),
            )
        )
        self.assertEqual(
            tuple(rule.rule_id for rule in forward.catalog.rules), expected_ids
        )
        for rule_id in expected_ids:
            self.assertEqual(
                forward.catalog.get(rule_id),
                reverse.catalog.get(rule_id),
            )
        self.assertEqual(len(forward.catalog.by_kind("range")), 1)
        self.assertEqual(len(forward.catalog.by_stage("pre_export")), 13)

    def test_diagnostic_order_is_independent_of_input_order(self):
        records = (
            make_record(rule_id="z_rule", validation_stage_id="future"),
            make_record(rule_id="a_rule", severity_id="future"),
            make_record(rule_id="m_rule", blocking="yes"),
        )
        forward = rules.build_validation_rule_catalog(records)
        reverse = rules.build_validation_rule_catalog(tuple(reversed(records)))
        self.assertEqual(
            tuple(item.as_dict() for item in forward.diagnostics),
            tuple(item.as_dict() for item in reverse.diagnostics),
        )
        self.assertEqual(
            tuple(item.rule_id for item in forward.diagnostics),
            tuple(sorted(item.rule_id for item in forward.diagnostics)),
        )

    def test_production_module_has_no_execution_or_io_surface(self):
        source = SCRIPT_PATH.read_text(encoding="utf-8")
        forbidden = (
            "openpyxl",
            "load_workbook",
            "import os",
            "import pathlib",
            "import random",
            "import socket",
            "eval(",
            "exec(",
            "getattr(",
            "setattr(",
            "ValidationStageRunner",
            "ValidationExecutionLedger",
        )
        for token in forbidden:
            with self.subTest(token=token):
                self.assertNotIn(token, source)
        self.assertEqual(len(rules.__all__), 23)


class TestWorkbookCorpus(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.headers, cls.inputs = read_active_rule_inputs()
        cls.result = rules.build_validation_rule_catalog(cls.inputs)

    def test_complete_active_corpus_builds_without_contract_errors(self):
        self.assertEqual(len(self.inputs), 308)
        self.assertTrue(self.result.is_valid, self.result.diagnostics)
        catalog = self.result.catalog
        self.assertEqual(len(catalog), 308)
        self.assertEqual(len(catalog.expression_rules), 249)
        self.assertEqual(len(catalog.structured_only_rules), 59)
        self.assertEqual(len({rule.rule_id for rule in catalog.rules}), 308)
        self.assertEqual(
            {rule.validation_kind_id for rule in catalog.rules},
            rules.KNOWN_VALIDATION_KINDS,
        )
        self.assertEqual(
            {rule.validation_stage_id for rule in catalog.rules},
            rules.KNOWN_VALIDATION_STAGES,
        )

    def test_workbook_headers_match_the_explicit_input_contract(self):
        self.assertEqual(set(self.headers), {"REGISTRY.xlsx", "CARDDATABASE.xlsx"})
        for workbook, headers in self.headers.items():
            with self.subTest(workbook=workbook):
                self.assertEqual(headers, rules.RULE_RECORD_FIELDS)

    def test_exact_kind_and_stage_distribution(self):
        self.assertEqual(
            dict(self.result.catalog.kind_counts),
            {
                "allowed_value": 5,
                "contract_field_invariant": 3,
                "contract_instance_consistency": 15,
                "cross_field_consistency": 1,
                "custom_expression": 57,
                "definition_invariant": 116,
                "hierarchy_integrity": 7,
                "normalized_uniqueness": 2,
                "qualified_reference_match": 7,
                "range": 10,
                "reference_integrity": 11,
                "target_group_membership": 1,
                "uniqueness": 73,
            },
        )
        self.assertEqual(
            dict(self.result.catalog.stage_counts),
            {
                "pre_export": 287,
                "production_export": 5,
                "runtime_load": 15,
                "runtime_package_build": 1,
            },
        )

    def test_reverse_workbook_input_is_catalog_identical(self):
        reverse = rules.build_validation_rule_catalog(tuple(reversed(self.inputs)))
        self.assertTrue(reverse.is_valid, reverse.diagnostics)
        self.assertEqual(self.result, reverse)

    def test_current_identifier_and_binding_inventory_matches_human_decision(self):
        identifier_references = 0
        binding_targets = 0
        current_node_references = 0
        current_node_bindings = 0
        for rule in self.result.catalog.expression_rules:
            for node in walk_ast(rule.condition_ast):
                if isinstance(node, rules.Identifier) and node.name == "current":
                    identifier_references += 1
                if isinstance(node, rules.Binding) and node.name == "current":
                    binding_targets += 1
                if isinstance(node, rules.Identifier) and node.name == "current_node":
                    current_node_references += 1
                if isinstance(node, rules.Binding) and node.name == "current_node":
                    current_node_bindings += 1
        self.assertEqual(identifier_references, 33)
        self.assertEqual(binding_targets, 0)
        self.assertEqual(current_node_references, 2)
        self.assertEqual(current_node_bindings, 1)
        local_current = tuple(
            rule.rule_id
            for rule in self.result.catalog.rules
            if "current" in rule.local_bindings
        )
        self.assertEqual(local_current, ())
        local_current_node = tuple(
            rule.rule_id
            for rule in self.result.catalog.rules
            if "current_node" in rule.local_bindings
        )
        self.assertEqual(
            local_current_node,
            ("val_ability_template_bindings_generated_node",),
        )
        self.assertEqual(
            sum("current" in rule.free_identifiers for rule in self.result.catalog.rules),
            20,
        )

    def test_structured_numeric_representation_is_preserved_as_openpyxl_float(self):
        range_rules = self.result.catalog.by_kind("range")
        self.assertEqual(len(range_rules), 10)
        self.assertTrue(
            all(type(rule.comparison_value) is float for rule in range_rules)
        )
        self.assertTrue(all(type(rule.minimum_value) is float for rule in range_rules))
        self.assertEqual({rule.comparison_value for rule in range_rules}, {1.0})
        self.assertEqual({rule.minimum_value for rule in range_rules}, {1.0})

    def test_event_types_hierarchy_rule_is_present_from_workbook(self):
        registry = load_workbook(WORKBOOKS[0], read_only=True, data_only=False)
        try:
            rows = registry["SCHEMA_FIELDS"].iter_rows(values_only=True)
            headers = tuple(next(rows))
            index = {name: position for position, name in enumerate(headers)}
            field_exists = any(
                row[index["table_id"]] == "event_types"
                and row[index["field_name"]] == "parent_event_type_id"
                for row in rows
            )
        finally:
            registry.close()
        self.assertTrue(field_exists)
        hierarchy_rules = tuple(
            rule
            for rule in self.result.catalog.by_kind("hierarchy_integrity")
            if rule.target_table_id == "event_types"
        )
        self.assertEqual(len(hierarchy_rules), 1)
        self.assertEqual(
            hierarchy_rules[0].rule_id,
            "val_event_types_hierarchy_acyclic",
        )


if __name__ == "__main__":
    unittest.main()
