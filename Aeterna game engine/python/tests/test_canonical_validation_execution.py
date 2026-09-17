import hashlib
import importlib.util
import subprocess
import sys
import unittest
from collections import Counter
from dataclasses import FrozenInstanceError, replace
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


PYTHON_ROOT = Path(__file__).resolve().parents[1]
REPOSITORY_ROOT = Path(__file__).resolve().parents[3]
MODULE_DIRECTORY = PYTHON_ROOT / "tools" / "canonical_export"
RULES_PATH = MODULE_DIRECTORY / "canonical_validation_rules.py"
CORE_PATH = MODULE_DIRECTORY / "canonical_validation_execution_core.py"
EXECUTION_PATH = MODULE_DIRECTORY / "canonical_validation_execution.py"
UNIQUENESS_PATH = MODULE_DIRECTORY / "canonical_validation_uniqueness.py"
QUALIFIED_REFERENCE_PATH = (
    MODULE_DIRECTORY / "canonical_validation_qualified_reference.py"
)
WORKBOOKS = (
    REPOSITORY_ROOT / "Aeterna dokumentációk" / "REGISTRY.xlsx",
    REPOSITORY_ROOT / "Aeterna dokumentációk" / "CARDDATABASE.xlsx",
)


def load_module(module_name, path):
    sys.path.insert(0, str(MODULE_DIRECTORY))
    try:
        spec = importlib.util.spec_from_file_location(module_name, path)
        module = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = module
        spec.loader.exec_module(module)
        return module
    finally:
        sys.path.remove(str(MODULE_DIRECTORY))


rules = load_module("canonical_validation_rules", RULES_PATH)
validation_expr = sys.modules["canonical_validation_expr"]
core = load_module("canonical_validation_execution_core", CORE_PATH)
uniqueness = load_module("canonical_validation_uniqueness", UNIQUENESS_PATH)
execution = load_module("canonical_validation_execution", EXECUTION_PATH)
qualified_reference = sys.modules["canonical_validation_qualified_reference"]


def make_rule(**overrides):
    record = {field: "#NULL" for field in rules.RULE_RECORD_FIELDS}
    record.update(
        {
            "validation_rule_id": "allowed_test_value",
            "rule_scope_id": "field",
            "target_table_id": "items",
            "target_field_id": "fld_items_value",
            "validation_kind_id": "allowed_value",
            "operator_id": "op_in",
            "comparison_value": "test_group",
            "severity_id": "critical",
            "blocking": True,
            "error_code": "TEST_VALUE_INVALID",
            "message": "Value must be an active canonical member.",
            "validation_stage_id": "pre_export",
            "status": "active",
            "source_id": "src_test",
            "source_ref": "ITEMS.value",
        }
    )
    record.update(overrides)
    result = rules.build_validation_rule_catalog((record,))
    if not result.is_valid:
        raise AssertionError(result.diagnostics)
    return result.catalog.rules[0]


def make_context(
    records=(),
    *,
    include_target=True,
    include_group=True,
    registry_records=None,
):
    tables = {
        "schema_tables": (
            {
                "table_id": "items",
                "primary_key": "item_id",
                "status": "active",
            },
        ),
        "schema_fields": (
            {
                "field_id": "fld_items_value",
                "table_id": "items",
                "field_name": "value",
                "status": "active",
            },
        ),
        "value_groups": (
            {"group_id": "test_group", "status": "active"},
        )
        if include_group
        else (),
        "value_registry": tuple(
            registry_records
            if registry_records is not None
            else (
                {
                    "registry_value_id": "test_group_alpha",
                    "group_id": "test_group",
                    "value_id": "alpha",
                    "lifecycle_status": "active",
                },
                {
                    "registry_value_id": "test_group_beta",
                    "group_id": "test_group",
                    "value_id": "beta",
                    "lifecycle_status": "active",
                },
            )
        ),
    }
    if include_target:
        tables["items"] = tuple(records)
    return execution.ValidationDataContext(tables)


def run(records=(), **context_options):
    return execution.execute_allowed_value_rule(
        make_rule(), make_context(records, **context_options)
    )


def make_range_rule(**overrides):
    record = {field: "#NULL" for field in rules.RULE_RECORD_FIELDS}
    record.update(
        {
            "validation_rule_id": "range_test_value",
            "rule_scope_id": "field",
            "target_table_id": "items",
            "target_field_id": "fld_items_order",
            "validation_kind_id": "range",
            "operator_id": "op_greater_than_or_equal",
            "comparison_value": 1.0,
            "minimum_value": 1.0,
            "severity_id": "critical",
            "blocking": True,
            "error_code": "TEST_ORDER_INVALID",
            "message": "Order must be a positive integer.",
            "validation_stage_id": "pre_export",
            "status": "active",
            "source_id": "src_test",
            "source_ref": "ITEMS.order",
        }
    )
    record.update(overrides)
    result = rules.build_validation_rule_catalog((record,))
    if not result.is_valid:
        raise AssertionError(result.diagnostics)
    return result.catalog.rules[0]


def make_range_context(records=(), *, include_target=True, field_overrides=None):
    field_schema = {
        "field_id": "fld_items_order",
        "table_id": "items",
        "field_name": "order",
        "data_type": "integer",
        "required_mode": "always",
        "nullable": False,
        "null_handling": "forbidden",
        "status": "active",
    }
    field_schema.update(field_overrides or {})
    tables = {
        "schema_tables": (
            {
                "table_id": "items",
                "primary_key": "item_id",
                "status": "active",
            },
        ),
        "schema_fields": (field_schema,),
    }
    if include_target:
        tables["items"] = tuple(records)
    return execution.ValidationDataContext(tables)


def run_range(records=(), **context_options):
    return execution.execute_range_rule(
        make_range_rule(), make_range_context(records, **context_options)
    )


def local_executor_tables(values, orders, *, allowed_values=None):
    return {
        "schema_tables": (
            {"table_id": "items", "primary_key": "item_id", "status": "active"},
        ),
        "schema_fields": (
            {
                "field_id": "fld_items_value",
                "table_id": "items",
                "field_name": "value",
                "status": "active",
            },
            {
                "field_id": "fld_items_order",
                "table_id": "items",
                "field_name": "order",
                "data_type": "integer",
                "required_mode": "always",
                "nullable": False,
                "null_handling": "forbidden",
                "status": "active",
            },
        ),
        "value_groups": ({"group_id": "test_group", "status": "active"},),
        "value_registry": tuple(
            {
                "registry_value_id": f"test_group_{value}",
                "group_id": "test_group",
                "value_id": value,
                "lifecycle_status": "active",
            }
            for value in sorted(set(allowed_values or values))
        ),
        "items": tuple(
            {"item_id": f"item-{index}", "value": value, "order": order}
            for index, (value, order) in enumerate(zip(values, orders), 1)
        ),
    }


def reverse_table_input(tables):
    return {
        table_id: tuple(reversed(records))
        for table_id, records in reversed(tuple(tables.items()))
    }


def local_executor_namespaced_context(*, reverse=False, bindings=None):
    namespaces = (
        ("registry", local_executor_tables(("alpha",), (1,))),
        ("carddatabase", local_executor_tables(("beta", "gamma"), (2, 3))),
    )
    if reverse:
        namespaces = tuple(
            (namespace, reverse_table_input(tables))
            for namespace, tables in reversed(namespaces)
        )
    return execution.ValidationDataContext(
        namespaced_tables=dict(namespaces),
        component_namespaces=bindings
        if bindings is not None
        else {
            "registry-component": "registry",
            "carddatabase-component": "carddatabase",
        },
    )


REFERENCE_POLICIES = {
    "table_identity_policy": "export_namespace_colon_table_id",
    "external_reference_identifier_policy": "namespace_colon_identifier",
}


def make_reference_rule(
    *,
    expression=None,
    target_field="reference_id",
    target_field_id="fld_items_reference_id",
    reference_table="references",
    reference_field_id="fld_references_reference_id",
    component="registry-component",
    **overrides,
):
    record = {field: "#NULL" for field in rules.RULE_RECORD_FIELDS}
    record.update(
        {
            "validation_rule_id": "reference_test_value",
            "rule_scope_id": "field",
            "target_table_id": "items",
            "target_field_id": target_field_id,
            "validation_kind_id": "reference_integrity",
            "condition_expression": expression or "#NULL",
            "reference_table_id": reference_table,
            "reference_field_id": reference_field_id,
            "severity_id": "critical",
            "blocking": True,
            "error_code": "TEST_REFERENCE_INVALID",
            "message": "Reference must resolve exactly once.",
            "validation_stage_id": "pre_export",
            "status": "active",
            "source_id": "src_test",
            "source_ref": f"ITEMS.{target_field}",
        }
    )
    record.update(overrides)
    result = rules.build_validation_rule_catalog(
        (rules.ValidationRuleInput(record, component),)
    )
    if not result.is_valid:
        raise AssertionError(result.diagnostics)
    return result.catalog.rules[0]


def reference_tables(
    *,
    namespace="registry",
    target_records=(),
    reference_records=None,
    target_field="reference_id",
    target_field_id="fld_items_reference_id",
    reference_field_id="fld_references_reference_id",
    guarded=False,
    include_target=True,
    include_reference=True,
    include_reference_schema=True,
    target_schema_overrides=None,
    reference_schema_overrides=None,
):
    target_schema = {
        "field_id": target_field_id,
        "table_id": "items",
        "field_name": target_field,
        "data_type": "string",
        "required_mode": "conditional" if guarded else "always",
        "nullable": guarded,
        "null_handling": "explicit_null" if guarded else "forbidden",
        "status": "active",
    }
    target_schema.update(target_schema_overrides or {})
    reference_schema = {
        "field_id": reference_field_id,
        "table_id": "references",
        "field_name": "reference_id",
        "data_type": "string",
        "required_mode": "always",
        "nullable": False,
        "null_handling": "forbidden",
        "status": "active",
    }
    reference_schema.update(reference_schema_overrides or {})
    schema_fields = [target_schema]
    if include_reference_schema:
        schema_fields.append(reference_schema)
    tables = {
        "schema_tables": (
            {"table_id": "items", "primary_key": "item_id", "status": "active"},
            {
                "table_id": "references",
                "primary_key": "reference_id",
                "status": "active",
            },
        ),
        "schema_fields": tuple(schema_fields),
    }
    if include_target:
        tables["items"] = tuple(target_records)
    if include_reference:
        tables["references"] = tuple(
            reference_records
            if reference_records is not None
            else ({"reference_id": "alpha", "status": "archived"},)
        )
    return tables


def make_reference_context(**options):
    namespace = options.pop("namespace", "registry")
    return execution.ValidationDataContext(
        namespaced_tables={namespace: reference_tables(namespace=namespace, **options)},
        component_namespaces={f"{namespace}-component": namespace},
        namespace_policies={namespace: REFERENCE_POLICIES},
    )


def run_reference(target_records=(), *, expression=None, **context_options):
    target_field = context_options.get("target_field", "reference_id")
    target_field_id = context_options.get(
        "target_field_id", "fld_items_reference_id"
    )
    rule = make_reference_rule(
        expression=expression,
        target_field=target_field,
        target_field_id=target_field_id,
    )
    context_options["target_records"] = target_records
    context_options["guarded"] = expression is not None
    return execution.execute_reference_integrity_rule(
        rule, make_reference_context(**context_options)
    )


def make_uniqueness_rule(
    *,
    component="registry-component",
    target_table="items",
    target_field="value",
    target_field_id="fld_items_value",
    **overrides,
):
    record = {field: "#NULL" for field in rules.RULE_RECORD_FIELDS}
    record.update(
        {
            "validation_rule_id": "uniqueness_test_value",
            "rule_scope_id": "table",
            "target_table_id": target_table,
            "target_field_id": target_field_id,
            "validation_kind_id": "uniqueness",
            "severity_id": "critical",
            "blocking": True,
            "error_code": "TEST_VALUE_DUPLICATE",
            "message": "Value must be unique.",
            "validation_stage_id": "pre_export",
            "status": "active",
            "source_id": "src_test",
            "source_ref": f"{target_table.upper()}.{target_field}",
        }
    )
    record.update(overrides)
    result = rules.build_validation_rule_catalog(
        (rules.ValidationRuleInput(record, component),)
    )
    if not result.is_valid:
        raise AssertionError(result.diagnostics)
    return result.catalog.rules[0]


def uniqueness_tables(
    records=(),
    *,
    target_table="items",
    target_field="value",
    target_field_id="fld_items_value",
    primary_key="item_id",
    include_target=True,
    include_table_schema=True,
    include_field_schema=True,
    field_overrides=None,
):
    field_schema = {
        "field_id": target_field_id,
        "table_id": target_table,
        "field_name": target_field,
        "data_type": "string",
        "required_mode": "always",
        "nullable": False,
        "null_handling": "forbidden",
        "status": "active",
    }
    field_schema.update(field_overrides or {})
    tables = {
        "schema_tables": (
            {
                "table_id": target_table,
                "primary_key": primary_key,
                "status": "active",
            },
        )
        if include_table_schema
        else (),
        "schema_fields": (field_schema,) if include_field_schema else (),
    }
    if include_target:
        tables[target_table] = tuple(records)
    return tables


def make_uniqueness_context(
    records=(),
    *,
    namespace="registry",
    component="registry-component",
    include_binding=True,
    **table_options,
):
    return execution.ValidationDataContext(
        namespaced_tables={
            namespace: uniqueness_tables(records, **table_options)
        },
        component_namespaces={component: namespace} if include_binding else {},
    )


def run_uniqueness(records=(), *, rule=None, **context_options):
    return execution.execute_uniqueness_rule(
        rule or make_uniqueness_rule(),
        make_uniqueness_context(records, **context_options),
    )


def composite_field(
    field_name,
    data_type="string",
    *,
    nullable=False,
    required_mode=None,
    null_handling=None,
    field_id=None,
):
    return {
        "field_id": field_id or f"fld_items_{field_name}",
        "table_id": "items",
        "field_name": field_name,
        "data_type": data_type,
        "required_mode": required_mode or ("conditional" if nullable else "always"),
        "nullable": nullable,
        "null_handling": null_handling or ("explicit_null" if nullable else "forbidden"),
        "status": "active",
    }


def make_composite_rule(field_names=("group_id", "sequence"), *, expression=None):
    target_field = field_names[-1] if field_names else "value"
    if expression is None:
        values = ",".join(f'"{field_name}"' for field_name in field_names)
        expression = f"unique_by([{values}])"
    return make_uniqueness_rule(
        target_field=target_field,
        target_field_id=f"fld_items_{target_field}",
        condition_expression=expression,
    )


def make_composite_context(
    records=(),
    *,
    fields=None,
    include_target=True,
    include_table_schema=True,
):
    fields = fields or (
        composite_field("group_id"),
        composite_field("sequence", "integer"),
    )
    tables = {
        "schema_tables": (
            {"table_id": "items", "primary_key": "item_id", "status": "active"},
        )
        if include_table_schema
        else (),
        "schema_fields": tuple(fields),
    }
    if include_target:
        tables["items"] = tuple(records)
    return execution.ValidationDataContext(
        namespaced_tables={"registry": tables},
        component_namespaces={"registry-component": "registry"},
    )


def run_composite(
    records=(), *, field_names=("group_id", "sequence"), fields=None, rule=None
):
    return execution.execute_uniqueness_rule(
        rule or make_composite_rule(field_names),
        make_composite_context(records, fields=fields),
    )


def make_non_null_rule(field_names=("value",), **overrides):
    target_field = field_names[-1] if field_names else "value"
    values = ",".join(f'"{field_name}"' for field_name in field_names)
    return make_uniqueness_rule(
        target_field=target_field,
        target_field_id=f"fld_items_{target_field}",
        condition_expression=f"unique_non_null_by([{values}])",
        **overrides,
    )


def run_non_null(records=(), *, field_names=("value",), fields=None, rule=None):
    fields = fields or tuple(
        composite_field(field_name, nullable=True) for field_name in field_names
    )
    return execution.execute_uniqueness_rule(
        rule or make_non_null_rule(field_names),
        make_composite_context(records, fields=fields),
    )


class TestExecutionCoreExtraction(unittest.TestCase):
    def test_uniqueness_module_is_standalone_and_facade_reexports_same_symbols(self):
        symbols = (
            "execute_uniqueness_rule",
            "VALIDATION_UNIQUENESS_DUPLICATE",
            "VALIDATION_UNIQUENESS_TARGET_FIELD_MISSING",
            "VALIDATION_UNIQUENESS_TARGET_RECORD_INVALID",
            "VALIDATION_UNIQUENESS_TARGET_TABLE_MISSING",
        )

        self.assertEqual(set(uniqueness.__all__), set(symbols))
        for symbol in symbols:
            with self.subTest(symbol=symbol):
                self.assertIs(getattr(execution, symbol), getattr(uniqueness, symbol))

        source = UNIQUENESS_PATH.read_text(encoding="utf-8")
        self.assertNotIn("import canonical_validation_execution", source)
        self.assertNotIn("from .canonical_validation_execution import", source)
        for token in ("openpyxl", "load_workbook", "pathlib", "Path(", "open("):
            with self.subTest(filesystem_token=token):
                self.assertNotIn(token, source)

    def test_execution_and_uniqueness_import_orders_work_in_fresh_processes(self):
        module_path = repr(str(MODULE_DIRECTORY))
        orders = (
            ("canonical_validation_uniqueness", "canonical_validation_execution"),
            ("canonical_validation_execution", "canonical_validation_uniqueness"),
        )

        for first, second in orders:
            script = (
                f"import sys; sys.path.insert(0, {module_path}); "
                f"import {first}; import {second}; "
                "from canonical_validation_execution import execute_uniqueness_rule; "
                "from canonical_validation_uniqueness import execute_uniqueness_rule as direct; "
                "assert execute_uniqueness_rule is direct"
            )
            with self.subTest(first=first):
                completed = subprocess.run(
                    [sys.executable, "-B", "-c", script],
                    cwd=PYTHON_ROOT,
                    capture_output=True,
                    text=True,
                    check=False,
                )
                self.assertEqual(
                    completed.returncode,
                    0,
                    completed.stdout + completed.stderr,
                )

    def test_public_facade_reexports_the_complete_existing_contract(self):
        names = (
            "ValidationOutcome",
            "ValidationExecutionDiagnostic",
            "ValidationExecutionResult",
            "ValidationDataContext",
            "VALIDATION_EXECUTOR_UNSUPPORTED",
            "VALIDATION_RULE_FAILED",
            "VALIDATION_CONTRACT_FIELD_INVARIANT_FAILED",
            "VALIDATION_CONTRACT_FIELD_TARGET_RECORD_INVALID",
            "VALIDATION_CONTRACT_FIELD_TARGET_TABLE_MISSING",
            "VALIDATION_CONTIGUOUS_EVALUATION_ERROR",
            "VALIDATION_CONTIGUOUS_ORDER_FAILED",
            "VALIDATION_CUSTOM_EXPRESSION_EVALUATION_ERROR",
            "VALIDATION_CUSTOM_EXPRESSION_FAILED",
            "VALIDATION_DEFINITION_INVARIANT_FAILED",
            "VALIDATION_DEFINITION_TARGET_RECORD_INVALID",
            "VALIDATION_DEFINITION_TARGET_TABLE_MISSING",
            "VALIDATION_HIERARCHY_INTEGRITY_FAILED",
            "VALIDATION_HIERARCHY_TARGET_RECORD_INVALID",
            "VALIDATION_HIERARCHY_TARGET_TABLE_MISSING",
            "VALIDATION_NORMALIZED_UNIQUENESS_CONFLICT",
            "VALIDATION_NORMALIZED_UNIQUENESS_EVALUATION_ERROR",
            "VALIDATION_CROSS_FIELD_CONSISTENCY_FAILED",
            "VALIDATION_REFERENCE_CONSTRAINT_EVALUATION_ERROR",
            "VALIDATION_TARGET_GROUP_MEMBERSHIP_FAILED",
            "VALIDATION_QUALIFIED_REFERENCE_INVARIANT_FAILED",
            "VALIDATION_QUALIFIED_REFERENCE_RESOLUTION_INVALID",
            "VALIDATION_QUALIFIED_REFERENCE_TARGET_RECORD_INVALID",
            "VALIDATION_QUALIFIED_REFERENCE_TARGET_TABLE_MISSING",
            "VALIDATION_UNIQUENESS_DUPLICATE",
            "VALIDATION_UNIQUENESS_TARGET_FIELD_MISSING",
            "VALIDATION_UNIQUENESS_TARGET_RECORD_INVALID",
            "VALIDATION_UNIQUENESS_TARGET_TABLE_MISSING",
            "execute_allowed_value_rule",
            "contiguous_ordering_capability_reason",
            "execute_contract_field_invariant_rule",
            "execute_contiguous_ordering_rule",
            "execute_scalar_string_pattern_custom_expression_rule",
            "execute_supported_custom_expression_rule",
            "execute_definition_invariant_rule",
            "execute_hierarchy_integrity_rule",
            "execute_normalized_uniqueness_rule",
            "execute_cross_field_consistency_rule",
            "execute_target_group_membership_rule",
            "execute_qualified_reference_match_rule",
            "execute_range_rule",
            "execute_reference_integrity_rule",
            "execute_uniqueness_rule",
            "scalar_string_pattern_capability_reason",
        )
        facade = __import__("canonical_validation_execution", fromlist=names)

        self.assertEqual(len(facade.__all__), 63)
        for name in names:
            with self.subTest(name=name):
                self.assertIn(name, facade.__all__)
                self.assertIs(getattr(facade, name), getattr(execution, name))
        for name in names[:6]:
            with self.subTest(core_symbol=name):
                self.assertIs(getattr(facade, name), getattr(core, name))
        for alias in ("CanonicalScalar", "FrozenRecord"):
            with self.subTest(compatibility_alias=alias):
                self.assertIs(getattr(facade, alias), getattr(core, alias))
                self.assertNotIn(alias, facade.__all__)

    def test_core_has_no_executor_rule_expression_or_io_dependency(self):
        source = CORE_PATH.read_text(encoding="utf-8")

        forbidden = (
            "from .canonical_validation_execution import",
            "from canonical_validation_execution import",
            "import canonical_validation_execution",
            "canonical_validation_rules",
            "canonical_validation_expr",
            "openpyxl",
            "load_workbook",
            "pathlib",
            "Path(",
            "open(",
        )
        for token in forbidden:
            with self.subTest(token=token):
                self.assertNotIn(token, source)

    def test_legacy_and_namespaced_resolvers_are_equivalent(self):
        tables = reference_tables(
            target_records=({"item_id": "item-1", "reference_id": "alpha"},)
        )
        legacy = execution.ValidationDataContext(tables)
        namespaced = execution.ValidationDataContext(
            namespaced_tables={"registry": tables},
            component_namespaces={"registry-component": "registry"},
            namespace_policies={"registry": REFERENCE_POLICIES},
        )

        legacy_field = core._resolve_field_schema(
            legacy, "items", field_id="fld_items_reference_id"
        )
        namespaced_field = core._resolve_field_schema(
            namespaced,
            "items",
            field_id="fld_items_reference_id",
            namespace="registry",
        )
        self.assertEqual(legacy_field, namespaced_field)
        self.assertEqual(
            core._resolve_field_schema(
                legacy, "references", field_name="reference_id"
            ),
            core._resolve_field_schema(
                namespaced,
                "references",
                field_name="reference_id",
                namespace="registry",
                allow_empty_resolved_name=True,
            ),
        )
        self.assertEqual(
            core._resolve_primary_key(legacy, "items"),
            core._resolve_primary_key(namespaced, "items", namespace="registry"),
        )
        self.assertEqual(
            core._resolve_table(legacy, "items"),
            core._resolve_table(namespaced, "items", namespace="registry"),
        )

    def test_existing_allowed_and_range_contexts_use_shared_resolvers(self):
        allowed = make_context()
        ranged = make_range_context()

        self.assertEqual(
            core._resolve_field_schema(
                allowed, "items", field_id="fld_items_value"
            )["field_name"],
            "value",
        )
        self.assertEqual(
            core._resolve_field_schema(
                ranged, "items", field_id="fld_items_order"
            )["field_name"],
            "order",
        )
        self.assertEqual(core._resolve_primary_key(allowed, "items"), "item_id")
        self.assertEqual(core._resolve_primary_key(ranged, "items"), "item_id")

    def test_namespace_collision_never_becomes_unqualified_first_match(self):
        context = table_identity_context(
            registry_values=("value_registry",),
            carddatabase_values=("cards",),
        )

        self.assertIsNone(core._resolve_table(context, "schema_tables"))
        self.assertIsNone(
            core._resolve_field_schema(
                context,
                "schema_tables",
                field_id="fld_registry_schema_tables_table_id",
            )
        )
        self.assertIsNone(core._resolve_primary_key(context, "schema_tables"))
        for namespace in ("registry", "carddatabase"):
            with self.subTest(namespace=namespace):
                field = core._resolve_field_schema(
                    context,
                    "schema_tables",
                    field_id=f"fld_{namespace}_schema_tables_table_id",
                    namespace=namespace,
                )
                self.assertEqual(field["field_name"], "table_id")
                self.assertEqual(
                    core._resolve_primary_key(
                        context, "schema_tables", namespace=namespace
                    ),
                    "table_id",
                )

    def test_missing_namespace_resolver_paths_fail_closed(self):
        context = table_identity_context()

        self.assertIsNone(
            core._resolve_table(context, "schema_tables", namespace="missing")
        )
        self.assertIsNone(
            core._resolve_field_schema(
                context,
                "schema_tables",
                field_id="fld_registry_schema_tables_table_id",
                namespace="missing",
            )
        )
        self.assertIsNone(
            core._resolve_primary_key(
                context, "schema_tables", namespace="missing"
            )
        )


class TestAllowedValuePass(unittest.TestCase):
    def test_one_valid_value_passes(self):
        result = run(({"item_id": "item-1", "value": "alpha"},))

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 1)
        self.assertEqual(result.violation_count, 0)
        self.assertEqual(result.diagnostics, ())

    def test_multiple_valid_values_pass(self):
        result = run(
            (
                {"item_id": "item-2", "value": "beta"},
                {"item_id": "item-1", "value": "alpha"},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 2)
        self.assertEqual(result.violation_count, 0)

    def test_empty_target_table_vacuously_passes(self):
        result = run(())

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 0)
        self.assertEqual(result.violation_count, 0)


class TestAllowedValueFail(unittest.TestCase):
    def test_one_invalid_value_fails_with_evidence(self):
        result = run(({"item_id": "item-1", "value": "unknown"},))

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.violation_count, 1)
        diagnostic = result.diagnostics[0]
        self.assertEqual(
            diagnostic.code, execution.VALIDATION_ALLOWED_VALUE_INVALID
        )
        self.assertEqual(diagnostic.rule_id, "allowed_test_value")
        self.assertEqual(diagnostic.table_id, "items")
        self.assertEqual(diagnostic.field_id, "fld_items_value")
        self.assertEqual(diagnostic.record_identity, "item-1")
        self.assertEqual(diagnostic.observed_value, "unknown")
        self.assertEqual(
            diagnostic.expected_contract,
            "active VALUE_REGISTRY.value_id in group test_group",
        )

    def test_multiple_invalid_values_each_fail(self):
        result = run(
            (
                {"item_id": "item-2", "value": "unknown-b"},
                {"item_id": "item-1", "value": "unknown-a"},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 2)
        self.assertEqual(result.violation_count, 2)
        self.assertEqual(
            tuple(item.record_identity for item in result.diagnostics),
            ("item-1", "item-2"),
        )

    def test_valid_and_invalid_mixture_fails_only_invalid_record(self):
        result = run(
            (
                {"item_id": "item-1", "value": "alpha"},
                {"item_id": "item-2", "value": "unknown"},
                {"item_id": "item-3", "value": "beta"},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 3)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(result.diagnostics[0].record_identity, "item-2")

    def test_missing_target_field_is_not_silently_skipped(self):
        result = run(({"item_id": "item-1"},))

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(result.diagnostics[0].code, execution.VALIDATION_RULE_FAILED)
        self.assertEqual(result.diagnostics[0].reason, "target_field_missing")
        self.assertEqual(result.diagnostics[0].record_identity, "item-1")

    def test_missing_target_table_is_not_a_pass(self):
        result = run(include_target=False)

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 0)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(result.diagnostics[0].code, execution.VALIDATION_RULE_FAILED)
        self.assertEqual(result.diagnostics[0].reason, "target_table_missing")

    def test_missing_allowed_group_is_blocking_failure(self):
        result = run(
            ({"item_id": "item-1", "value": "alpha"},),
            include_group=False,
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 0)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(
            result.diagnostics[0].code,
            execution.VALIDATION_ALLOWED_VALUE_GROUP_MISSING,
        )

    def test_null_target_value_is_invalid(self):
        result = run(({"item_id": "item-1", "value": None},))

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(
            result.diagnostics[0].code,
            execution.VALIDATION_ALLOWED_VALUE_INVALID,
        )
        self.assertIsNone(result.diagnostics[0].observed_value)

    def test_alias_like_or_case_changed_token_is_not_normalized(self):
        for value in ("Alpha", " alpha ", "ALPHA", "legacy-alpha"):
            with self.subTest(value=value):
                result = run(({"item_id": "item-1", "value": value},))
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(result.diagnostics[0].observed_value, value)

    def test_inactive_or_deprecated_registry_value_is_not_allowed(self):
        for lifecycle in ("inactive", "deprecated"):
            with self.subTest(lifecycle=lifecycle):
                result = run(
                    ({"item_id": "item-1", "value": "retired"},),
                    registry_records=(
                        {
                            "registry_value_id": "test_group_alpha",
                            "group_id": "test_group",
                            "value_id": "alpha",
                            "lifecycle_status": "active",
                        },
                        {
                            "registry_value_id": "test_group_retired",
                            "group_id": "test_group",
                            "value_id": "retired",
                            "lifecycle_status": lifecycle,
                        },
                    ),
                )
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)


class TestAllowedValueContract(unittest.TestCase):
    def test_non_allowed_value_rule_is_unsupported(self):
        result = execution.execute_allowed_value_rule(
            replace(make_rule(), validation_kind_id="range"),
            make_context(()),
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
        self.assertEqual(result.violation_count, 0)
        self.assertEqual(
            result.diagnostics[0].code, execution.VALIDATION_EXECUTOR_UNSUPPORTED
        )
        self.assertEqual(
            result.diagnostics[0].reason, "validation_kind_not_allowed_value"
        )

    def test_unsupported_operator_is_reported_without_fallback(self):
        result = execution.execute_allowed_value_rule(
            replace(make_rule(), operator_id="op_equal"),
            make_context(()),
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
        self.assertEqual(result.diagnostics[0].reason, "operator_not_op_in")

    def test_unsupported_shape_is_reported_without_fallback(self):
        result = execution.execute_allowed_value_rule(
            replace(make_rule(), minimum_value=1),
            make_context(()),
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
        self.assertEqual(
            result.diagnostics[0].reason, "structured_shape_not_supported"
        )

    def test_malformed_data_context_is_rejected_at_boundary(self):
        cases = (
            None,
            [],
            {"items": "not-a-record-sequence"},
            {"items": ("not-a-record",)},
            {"items": ({"value": object()},)},
            {"": ()},
        )
        for value in cases:
            with self.subTest(value=value):
                with self.assertRaises(ValueError):
                    execution.ValidationDataContext(value)

    def test_data_context_copies_records_and_exposes_immutable_mappings(self):
        source = {"item_id": "item-1", "value": "alpha"}
        context = make_context((source,))
        source["value"] = "changed"

        self.assertEqual(context.table("items")[0]["value"], "alpha")
        with self.assertRaises(TypeError):
            context.table("items")[0]["value"] = "changed"
        with self.assertRaises(FrozenInstanceError):
            context._tables = ()

    def test_production_module_has_no_workbook_or_filesystem_dependency(self):
        source = EXECUTION_PATH.read_text(encoding="utf-8")
        qualified_source = QUALIFIED_REFERENCE_PATH.read_text(encoding="utf-8")

        for token in ("openpyxl", "load_workbook", "pathlib", "Path("):
            with self.subTest(token=token):
                self.assertNotIn(token, source)
                self.assertNotIn(token, qualified_source)
        self.assertNotIn(
            "from .canonical_validation_execution import", qualified_source
        )
        self.assertNotIn(
            "from canonical_validation_execution import", qualified_source
        )
        self.assertEqual(len(execution.__all__), 63)


class TestAllowedValueDeterminism(unittest.TestCase):
    def test_reordered_input_has_identical_result_and_diagnostics(self):
        records = (
            {"item_id": "item-3", "value": "unknown-c"},
            {"item_id": "item-1", "value": "alpha"},
            {"item_id": "item-2", "value": "unknown-b"},
        )

        forward = run(records)
        reverse = run(tuple(reversed(records)))

        self.assertEqual(forward, reverse)
        self.assertEqual(
            tuple(item.record_identity for item in forward.diagnostics),
            ("item-2", "item-3"),
        )

    def test_missing_primary_key_uses_stable_record_content_not_input_index(self):
        records = (
            {"value": "unknown-b", "other": 2},
            {"value": "unknown-a", "other": 1},
        )

        forward = run(records)
        reverse = run(tuple(reversed(records)))

        self.assertEqual(forward, reverse)
        self.assertTrue(
            all(
                item.record_identity.startswith("canonical-record:")
                for item in forward.diagnostics
            )
        )


class TestRangePass(unittest.TestCase):
    def test_exact_inclusive_boundary_passes(self):
        result = run_range(({"item_id": "item-1", "order": 1},))

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 1)
        self.assertEqual(result.violation_count, 0)
        self.assertEqual(result.diagnostics, ())

    def test_values_above_boundary_and_multiple_records_pass(self):
        result = run_range(
            (
                {"item_id": "item-2", "order": 10},
                {"item_id": "item-1", "order": 2},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 2)
        self.assertEqual(result.violation_count, 0)

    def test_empty_existing_target_table_vacuously_passes(self):
        result = run_range(())

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 0)
        self.assertEqual(result.violation_count, 0)


class TestRangeFail(unittest.TestCase):
    def test_value_below_boundary_fails_with_evidence(self):
        result = run_range(({"item_id": "item-1", "order": 0},))

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.violation_count, 1)
        diagnostic = result.diagnostics[0]
        self.assertEqual(diagnostic.code, execution.VALIDATION_RANGE_INVALID)
        self.assertEqual(diagnostic.rule_id, "range_test_value")
        self.assertEqual(diagnostic.table_id, "items")
        self.assertEqual(diagnostic.field_id, "fld_items_order")
        self.assertEqual(diagnostic.record_identity, "item-1")
        self.assertEqual(diagnostic.observed_value, 0)
        self.assertEqual(
            diagnostic.expected_contract, "canonical integer value >= 1"
        )

    def test_multiple_invalid_values_each_fail(self):
        result = run_range(
            (
                {"item_id": "item-2", "order": -1},
                {"item_id": "item-1", "order": 0},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 2)
        self.assertEqual(result.violation_count, 2)
        self.assertEqual(
            tuple(item.record_identity for item in result.diagnostics),
            ("item-1", "item-2"),
        )

    def test_valid_and_invalid_mixture_reports_only_invalid_record(self):
        result = run_range(
            (
                {"item_id": "item-1", "order": 1},
                {"item_id": "item-2", "order": 0},
                {"item_id": "item-3", "order": 3},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 3)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(result.diagnostics[0].record_identity, "item-2")

    def test_null_is_a_type_failure_not_not_applicable(self):
        result = run_range(({"item_id": "item-1", "order": None},))

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(
            result.diagnostics[0].code,
            execution.VALIDATION_RANGE_VALUE_TYPE_INVALID,
        )
        self.assertIsNone(result.diagnostics[0].observed_value)

    def test_noncanonical_integer_types_are_not_coerced(self):
        for value in (1.0, 1.5, "1", True, False):
            with self.subTest(value=value, value_type=type(value).__name__):
                result = run_range(({"item_id": "item-1", "order": value},))
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(
                    result.diagnostics[0].code,
                    execution.VALIDATION_RANGE_VALUE_TYPE_INVALID,
                )
                self.assertIs(result.diagnostics[0].observed_value, value)

    def test_missing_target_table_is_not_a_pass(self):
        result = run_range(include_target=False)

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 0)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(result.diagnostics[0].code, execution.VALIDATION_RULE_FAILED)
        self.assertEqual(result.diagnostics[0].reason, "target_table_missing")

    def test_missing_target_field_is_not_silently_skipped(self):
        result = run_range(({"item_id": "item-1"},))

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(result.diagnostics[0].code, execution.VALIDATION_RULE_FAILED)
        self.assertEqual(result.diagnostics[0].reason, "target_field_missing")
        self.assertEqual(result.diagnostics[0].record_identity, "item-1")


class TestRangeContract(unittest.TestCase):
    def test_integral_float_rule_literals_are_explicit_storage_artifacts(self):
        rule = make_range_rule(comparison_value=1.0, minimum_value=1.0)
        result = execution.execute_range_rule(
            rule,
            make_range_context(({"item_id": "item-1", "order": 1},)),
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)

    def test_non_range_rule_is_unsupported(self):
        result = execution.execute_range_rule(
            replace(make_range_rule(), validation_kind_id="allowed_value"),
            make_range_context(()),
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
        self.assertEqual(result.violation_count, 0)
        self.assertEqual(
            result.diagnostics[0].code, execution.VALIDATION_EXECUTOR_UNSUPPORTED
        )
        self.assertEqual(result.diagnostics[0].reason, "validation_kind_not_range")

    def test_unsupported_operator_is_reported_without_fallback(self):
        result = execution.execute_range_rule(
            replace(make_range_rule(), operator_id="op_less_than_or_equal"),
            make_range_context(()),
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
        self.assertEqual(
            result.diagnostics[0].reason, "operator_not_greater_than_or_equal"
        )

    def test_upper_bound_shape_is_not_invented(self):
        result = execution.execute_range_rule(
            replace(make_range_rule(), maximum_value=10.0),
            make_range_context(()),
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
        self.assertEqual(
            result.diagnostics[0].reason, "structured_shape_not_supported"
        )

    def test_comparison_and_minimum_must_declare_same_boundary(self):
        result = execution.execute_range_rule(
            replace(make_range_rule(), comparison_value=2.0, minimum_value=1.0),
            make_range_context(()),
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
        self.assertEqual(
            result.diagnostics[0].reason, "comparison_minimum_mismatch"
        )

    def test_non_required_integer_target_schema_is_unsupported(self):
        for override in (
            {"data_type": "number"},
            {"required_mode": "conditional"},
            {"nullable": True},
            {"null_handling": "explicit_null"},
        ):
            with self.subTest(override=override):
                result = execution.execute_range_rule(
                    make_range_rule(),
                    make_range_context((), field_overrides=override),
                )
                self.assertEqual(
                    result.outcome, execution.ValidationOutcome.UNSUPPORTED
                )
                self.assertEqual(
                    result.diagnostics[0].reason,
                    "target_schema_not_required_integer",
                )


class TestRangeDeterminism(unittest.TestCase):
    def test_reordered_input_has_identical_result_and_diagnostics(self):
        records = (
            {"item_id": "item-3", "order": -2},
            {"item_id": "item-1", "order": 1},
            {"item_id": "item-2", "order": 0},
        )

        forward = run_range(records)
        reverse = run_range(tuple(reversed(records)))

        self.assertEqual(forward, reverse)
        self.assertEqual(
            tuple(item.record_identity for item in forward.diagnostics),
            ("item-2", "item-3"),
        )


class TestLocalExecutorNamespaces(unittest.TestCase):
    def test_allowed_value_routes_each_component_to_its_local_namespace(self):
        context = local_executor_namespaced_context()
        cases = (
            ("registry-component", 1),
            ("carddatabase-component", 2),
        )

        for component, expected_count in cases:
            with self.subTest(component=component):
                rule = replace(make_rule(), component_identity=component)
                result = execution.execute_allowed_value_rule(rule, context)
                self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
                self.assertEqual(result.evaluated_record_count, expected_count)
                self.assertEqual(result.violation_count, 0)
                self.assertEqual(result.diagnostics, ())

    def test_range_routes_each_component_to_its_local_namespace(self):
        context = local_executor_namespaced_context()
        cases = (
            ("registry-component", 1),
            ("carddatabase-component", 2),
        )

        for component, expected_count in cases:
            with self.subTest(component=component):
                rule = replace(make_range_rule(), component_identity=component)
                result = execution.execute_range_rule(rule, context)
                self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
                self.assertEqual(result.evaluated_record_count, expected_count)
                self.assertEqual(result.violation_count, 0)
                self.assertEqual(result.diagnostics, ())

    def test_same_table_collision_is_independent_of_namespace_input_order(self):
        rules_and_executors = (
            (
                replace(make_rule(), component_identity="registry-component"),
                execution.execute_allowed_value_rule,
            ),
            (
                replace(make_rule(), component_identity="carddatabase-component"),
                execution.execute_allowed_value_rule,
            ),
            (
                replace(
                    make_range_rule(), component_identity="registry-component"
                ),
                execution.execute_range_rule,
            ),
            (
                replace(
                    make_range_rule(), component_identity="carddatabase-component"
                ),
                execution.execute_range_rule,
            ),
        )

        def execute_all(context):
            return tuple(executor(rule, context) for rule, executor in rules_and_executors)

        forward = execute_all(local_executor_namespaced_context())
        reverse = execute_all(local_executor_namespaced_context(reverse=True))

        self.assertEqual(forward, reverse)
        self.assertEqual(
            tuple(result.evaluated_record_count for result in forward),
            (1, 2, 1, 2),
        )

    def test_missing_component_binding_fails_closed_without_lookup(self):
        context = local_executor_namespaced_context(bindings={})
        rules_and_executors = (
            (
                replace(make_rule(), component_identity="missing-component"),
                execution.execute_allowed_value_rule,
            ),
            (
                replace(
                    make_range_rule(), component_identity="missing-component"
                ),
                execution.execute_range_rule,
            ),
        )

        for rule, executor in rules_and_executors:
            with self.subTest(kind=rule.validation_kind_id):
                result = executor(rule, context)
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(result.evaluated_record_count, 0)
                self.assertEqual(result.violation_count, 1)
                self.assertEqual(len(result.diagnostics), 1)
                self.assertEqual(
                    result.diagnostics[0].code, execution.VALIDATION_RULE_FAILED
                )
                self.assertEqual(
                    result.diagnostics[0].reason, "source_namespace_missing"
                )

    def test_wrong_binding_does_not_fall_back_to_another_namespace(self):
        context = execution.ValidationDataContext(
            namespaced_tables={
                "registry": local_executor_tables(("alpha",), (1,)),
                "carddatabase": local_executor_tables(
                    ("invalid",), (0,), allowed_values=("different",)
                ),
            },
            component_namespaces={"registry-component": "carddatabase"},
        )
        rules_and_executors = (
            (
                replace(make_rule(), component_identity="registry-component"),
                execution.execute_allowed_value_rule,
                execution.VALIDATION_ALLOWED_VALUE_INVALID,
            ),
            (
                replace(
                    make_range_rule(), component_identity="registry-component"
                ),
                execution.execute_range_rule,
                execution.VALIDATION_RANGE_INVALID,
            ),
        )

        for rule, executor, diagnostic_code in rules_and_executors:
            with self.subTest(kind=rule.validation_kind_id):
                result = executor(rule, context)
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(result.evaluated_record_count, 1)
                self.assertEqual(result.violation_count, 1)
                self.assertEqual(result.diagnostics[0].code, diagnostic_code)

        with self.assertRaises(ValueError):
            execution.ValidationDataContext(
                namespaced_tables={
                    "registry": local_executor_tables(("alpha",), (1,))
                },
                component_namespaces={"registry-component": "missing"},
            )

    def test_legacy_context_results_remain_compatible(self):
        namespaced = local_executor_namespaced_context()
        allowed_rule = replace(
            make_rule(), component_identity="registry-component"
        )
        range_rule = replace(
            make_range_rule(), component_identity="registry-component"
        )

        self.assertEqual(
            execution.execute_allowed_value_rule(allowed_rule, namespaced),
            run(({"item_id": "item-1", "value": "alpha"},)),
        )
        self.assertEqual(
            execution.execute_range_rule(range_rule, namespaced),
            run_range(({"item_id": "item-1", "order": 1},)),
        )


class TestUniquenessPass(unittest.TestCase):
    def test_empty_one_and_several_distinct_records_pass(self):
        cases = (
            (),
            ({"item_id": "item-1", "value": "alpha"},),
            (
                {"item_id": "item-2", "value": "beta"},
                {"item_id": "item-1", "value": "alpha"},
                {"item_id": "item-3", "value": "gamma"},
            ),
        )
        for records in cases:
            with self.subTest(records=records):
                result = run_uniqueness(records)
                self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
                self.assertEqual(result.evaluated_record_count, len(records))
                self.assertEqual(result.violation_count, 0)
                self.assertEqual(result.diagnostics, ())

    def test_case_whitespace_and_unicode_distinctions_are_exact(self):
        values = ("A", "a", "foo", " foo", "\u00e9", "e\u0301")
        records = tuple(
            {"item_id": f"item-{index}", "value": value}
            for index, value in enumerate(values, start=1)
        )

        result = run_uniqueness(records)

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, len(values))
        self.assertEqual(result.diagnostics, ())

    def test_mixed_lifecycle_and_export_flags_are_not_filtered(self):
        records = (
            {"item_id": "item-1", "value": "alpha", "status": "active"},
            {"item_id": "item-2", "value": "beta", "status": "superseded"},
            {"item_id": "item-3", "value": "gamma", "status": "archived"},
            {"item_id": "item-4", "value": "delta", "export_enabled": False},
        )

        result = run_uniqueness(records)

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 4)

    def test_primary_key_and_non_primary_targets_both_execute(self):
        non_primary = run_uniqueness(
            (
                {"item_id": "item-1", "value": "alpha"},
                {"item_id": "item-2", "value": "beta"},
            )
        )
        primary_rule = make_uniqueness_rule(
            target_field="item_id", target_field_id="fld_items_item_id"
        )
        primary_context = make_uniqueness_context(
            ({"item_id": "item-1"}, {"item_id": "item-2"}),
            target_field="item_id",
            target_field_id="fld_items_item_id",
        )
        primary = execution.execute_uniqueness_rule(primary_rule, primary_context)

        self.assertEqual(non_primary.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(primary.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual((non_primary.evaluated_record_count, primary.evaluated_record_count), (2, 2))


class TestUniquenessDuplicates(unittest.TestCase):
    def test_simple_duplicate_is_one_group_diagnostic(self):
        result = run_uniqueness(
            (
                {"item_id": "item-2", "value": "A"},
                {"item_id": "item-1", "value": "A"},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 2)
        self.assertEqual(result.violation_count, 1)
        diagnostic = result.diagnostics[0]
        self.assertEqual(diagnostic.code, execution.VALIDATION_UNIQUENESS_DUPLICATE)
        self.assertEqual(diagnostic.rule_id, "uniqueness_test_value")
        self.assertEqual(diagnostic.table_id, "items")
        self.assertEqual(diagnostic.field_id, "fld_items_value")
        self.assertEqual(diagnostic.observed_value, "A")
        self.assertEqual(diagnostic.related_record_identities, ("item-1", "item-2"))
        self.assertIn("registry:items.value", diagnostic.expected_contract)

    def test_triple_duplicate_is_one_group_with_three_identities(self):
        result = run_uniqueness(
            tuple(
                {"item_id": f"item-{index}", "value": "A"}
                for index in (3, 1, 2)
            )
        )

        self.assertEqual(result.violation_count, 1)
        self.assertEqual(
            result.diagnostics[0].related_record_identities,
            ("item-1", "item-2", "item-3"),
        )

    def test_two_duplicate_groups_produce_two_diagnostics(self):
        records = (
            {"item_id": "item-5", "value": "C"},
            {"item_id": "item-4", "value": "B"},
            {"item_id": "item-2", "value": "A"},
            {"item_id": "item-3", "value": "B"},
            {"item_id": "item-1", "value": "A"},
        )

        result = run_uniqueness(records)

        self.assertEqual(result.violation_count, 2)
        self.assertEqual(
            tuple(item.observed_value for item in result.diagnostics),
            ("A", "B"),
        )
        self.assertEqual(
            tuple(item.related_record_identities for item in result.diagnostics),
            (("item-1", "item-2"), ("item-3", "item-4")),
        )

    def test_duplicate_non_primary_field_is_not_treated_as_primary_key_check(self):
        result = run_uniqueness(
            (
                {"item_id": "different-1", "value": "duplicate"},
                {"item_id": "different-2", "value": "duplicate"},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.violation_count, 1)

    def test_primary_key_target_is_not_skipped(self):
        rule = make_uniqueness_rule(
            target_field="item_id", target_field_id="fld_items_item_id"
        )
        context = make_uniqueness_context(
            ({"item_id": "duplicate"}, {"item_id": "duplicate"}),
            target_field="item_id",
            target_field_id="fld_items_item_id",
        )

        result = execution.execute_uniqueness_rule(rule, context)

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.diagnostics[0].observed_value, "duplicate")
        self.assertEqual(
            result.diagnostics[0].related_record_identities,
            ("duplicate", "duplicate"),
        )

    def test_duplicate_crosses_lifecycle_statuses(self):
        result = run_uniqueness(
            (
                {"item_id": "item-1", "value": "A", "status": "active"},
                {"item_id": "item-2", "value": "A", "status": "archived"},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.violation_count, 1)

    def test_identical_precomposed_unicode_is_duplicate(self):
        result = run_uniqueness(
            (
                {"item_id": "item-1", "value": "\u00e9"},
                {"item_id": "item-2", "value": "\u00e9"},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.diagnostics[0].observed_value, "\u00e9")


class TestUniquenessMalformedTargets(unittest.TestCase):
    def test_missing_table_fails_closed(self):
        result = run_uniqueness((), include_target=False)

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 0)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(
            result.diagnostics[0].code,
            execution.VALIDATION_UNIQUENESS_TARGET_TABLE_MISSING,
        )
        self.assertEqual(result.diagnostics[0].reason, "target_table_missing")

    def test_missing_table_schema_fails_closed(self):
        result = run_uniqueness((), include_table_schema=False)

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(
            result.diagnostics[0].code,
            execution.VALIDATION_UNIQUENESS_TARGET_TABLE_MISSING,
        )
        self.assertEqual(
            result.diagnostics[0].reason, "target_table_schema_unresolved"
        )

    def test_missing_and_malformed_field_schema_fail_closed(self):
        missing = run_uniqueness((), include_field_schema=False)
        malformed = run_uniqueness((), field_overrides={"field_name": ""})

        for result in (missing, malformed):
            self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
            self.assertEqual(
                result.diagnostics[0].code,
                execution.VALIDATION_UNIQUENESS_TARGET_FIELD_MISSING,
            )
        self.assertEqual(missing.diagnostics[0].reason, "target_field_schema_unresolved")
        self.assertEqual(malformed.diagnostics[0].reason, "target_field_schema_malformed")

    def test_missing_null_blank_integer_and_boolean_values_fail_per_record(self):
        cases = (
            ({"item_id": "item-1"}, "target_field_missing"),
            ({"item_id": "item-1", "value": None}, "target_value_null_forbidden"),
            ({"item_id": "item-1", "value": ""}, "target_value_blank"),
            ({"item_id": "item-1", "value": "   "}, "target_value_blank"),
            ({"item_id": "item-1", "value": 1}, "target_value_not_string"),
            ({"item_id": "item-1", "value": True}, "target_value_not_string"),
        )
        for record, reason in cases:
            with self.subTest(record=record):
                result = run_uniqueness((record,))
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(result.evaluated_record_count, 1)
                self.assertEqual(result.violation_count, 1)
                self.assertEqual(
                    result.diagnostics[0].code,
                    execution.VALIDATION_UNIQUENESS_TARGET_RECORD_INVALID,
                )
                self.assertEqual(result.diagnostics[0].reason, reason)

    def test_missing_record_identity_fails_with_stable_fallback_identity(self):
        result = run_uniqueness(({"value": "alpha"},))

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.diagnostics[0].reason, "target_identity_invalid")
        self.assertTrue(
            result.diagnostics[0].record_identity.startswith("canonical-record:")
        )

    def test_list_and_object_values_are_rejected_at_context_boundary(self):
        for value in ([], {}):
            with self.subTest(value=value):
                with self.assertRaises(ValueError):
                    make_uniqueness_context(
                        ({"item_id": "item-1", "value": value},)
                    )

    def test_missing_source_namespace_fails_closed(self):
        result = run_uniqueness((), include_binding=False)

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(
            result.diagnostics[0].code,
            execution.VALIDATION_UNIQUENESS_TARGET_TABLE_MISSING,
        )
        self.assertEqual(result.diagnostics[0].reason, "source_namespace_missing")

    def test_ambiguous_namespace_provenance_is_rejected_at_context_boundary(self):
        tables = uniqueness_tables(())
        with self.assertRaises(ValueError):
            execution.ValidationDataContext(
                namespaced_tables={"registry": tables},
                component_namespaces={
                    "component-a": "registry",
                    "component-b": "registry",
                },
            )

    def test_nullable_and_non_string_schemas_are_unsupported(self):
        overrides = (
            {
                "required_mode": "conditional",
                "nullable": True,
                "null_handling": "explicit_null",
            },
            {"data_type": "integer"},
        )
        for field_override in overrides:
            with self.subTest(field_override=field_override):
                result = run_uniqueness((), field_overrides=field_override)
                self.assertEqual(
                    result.outcome, execution.ValidationOutcome.UNSUPPORTED
                )
                self.assertEqual(result.violation_count, 0)
                self.assertEqual(
                    result.diagnostics[0].reason,
                    "target_schema_not_required_string",
                )

    def test_unsupported_rule_shapes_have_no_fallback(self):
        base = make_uniqueness_rule()
        cases = (
            (replace(base, validation_kind_id="range"), "validation_kind_not_uniqueness"),
            (replace(base, rule_scope_id="field"), "rule_scope_not_table"),
            (replace(base, validation_stage_id="runtime_load"), "validation_stage_not_pre_export"),
            (replace(base, target_field_id=None), "target_field_missing"),
            (replace(base, operator_id="op_equal"), "structured_shape_not_supported"),
        )
        context = make_uniqueness_context(())
        for rule, reason in cases:
            with self.subTest(reason=reason):
                result = execution.execute_uniqueness_rule(rule, context)
                self.assertEqual(
                    result.outcome, execution.ValidationOutcome.UNSUPPORTED
                )
                self.assertEqual(result.violation_count, 0)
                self.assertEqual(result.diagnostics[0].reason, reason)

    def test_unique_non_null_by_is_the_third_supported_family_shape(self):
        result = run_non_null(())

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 0)
        self.assertEqual(result.diagnostics, ())


class TestCompositeUniquenessAstContract(unittest.TestCase):
    def test_valid_one_two_three_and_four_field_shapes_execute(self):
        for width in range(1, 5):
            field_names = tuple(f"field_{index}" for index in range(width))
            fields = tuple(composite_field(name) for name in field_names)
            with self.subTest(width=width):
                result = run_composite((), field_names=field_names, fields=fields)
                self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
                self.assertEqual(result.evaluated_record_count, 0)

    def test_invalid_ast_shapes_are_unsupported_without_partial_execution(self):
        literal = validation_expr.Literal("string", "group_id")
        field_list = validation_expr.ListLiteral((literal,))
        base = make_composite_rule()
        cases = (
            make_composite_rule((), expression="unique_by([])"),
            make_composite_rule(("group_id", "group_id")),
            make_uniqueness_rule(
                target_field="group_id",
                target_field_id="fld_items_group_id",
                condition_expression='unique_by(["group_id",1])',
            ),
            replace(
                base,
                condition_ast=validation_expr.Call(
                    "unique_by", (field_list, field_list)
                ),
            ),
            replace(
                base,
                condition_ast=validation_expr.Call(
                    "all_unique", (field_list,)
                ),
            ),
            replace(
                base,
                condition_ast=validation_expr.Call(
                    "unique_by",
                    (
                        validation_expr.ListLiteral(
                            (validation_expr.Call("trim", (literal,)),)
                        ),
                    ),
                ),
            ),
            replace(
                base,
                condition_ast=validation_expr.Sequence((base.condition_ast,)),
            ),
        )
        context = make_composite_context(())
        for rule in cases:
            with self.subTest(ast=rule.condition_ast):
                result = execution.execute_uniqueness_rule(rule, context)
                self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
                self.assertEqual(result.evaluated_record_count, 0)
                self.assertEqual(result.violation_count, 0)
                self.assertEqual(
                    result.diagnostics[0].code,
                    execution.VALIDATION_EXECUTOR_UNSUPPORTED,
                )

    def test_unresolved_ambiguous_and_unsupported_fields_fail_closed(self):
        unresolved = make_uniqueness_rule(
            target_field="group_id",
            target_field_id="fld_items_group_id",
            condition_expression='unique_by(["group_id","missing"])',
        )
        ambiguous = make_composite_rule()
        unsupported = make_composite_rule()
        cases = (
            (
                unresolved,
                (composite_field("group_id"),),
                "unique_by_field_unresolved:missing",
            ),
            (
                ambiguous,
                (
                    composite_field("group_id", field_id="fld_items_group_a"),
                    composite_field("group_id", field_id="fld_items_group_b"),
                    composite_field("sequence", "integer"),
                ),
                "unique_by_field_unresolved:group_id",
            ),
            (
                unsupported,
                (
                    composite_field("group_id", "boolean"),
                    composite_field("sequence", "integer"),
                ),
                "unique_by_field_schema_unsupported:group_id",
            ),
        )
        for rule, fields, reason in cases:
            with self.subTest(reason=reason):
                result = execution.execute_uniqueness_rule(
                    rule, make_composite_context((), fields=fields)
                )
                self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
                self.assertEqual(result.diagnostics[0].reason, reason)

    def test_missing_composite_target_table_is_a_failure(self):
        result = execution.execute_uniqueness_rule(
            make_composite_rule(), make_composite_context((), include_target=False)
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 0)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(
            result.diagnostics[0].code,
            execution.VALIDATION_UNIQUENESS_TARGET_TABLE_MISSING,
        )


class TestCompositeUniquenessPass(unittest.TestCase):
    def test_distinct_tuples_by_each_constituent_pass(self):
        result = run_composite(
            (
                {"item_id": "item-1", "group_id": "A", "sequence": 1},
                {"item_id": "item-2", "group_id": "B", "sequence": 1},
                {"item_id": "item-3", "group_id": "A", "sequence": 2},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 3)
        self.assertEqual(result.diagnostics, ())

    def test_string_case_whitespace_and_unicode_are_not_normalized(self):
        values = ("A", "a", " A", "\u00e9", "e\u0301")
        records = tuple(
            {"item_id": f"item-{index}", "group_id": value, "sequence": 1}
            for index, value in enumerate(values, 1)
        )

        result = run_composite(records)

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, len(records))

    def test_language_tag_is_exact_and_keeps_its_schema_type(self):
        fields = (
            composite_field("entity_id"),
            composite_field("language", "language_tag"),
        )
        result = run_composite(
            (
                {"item_id": "item-1", "entity_id": "card", "language": "hu"},
                {"item_id": "item-2", "entity_id": "card", "language": "HU"},
            ),
            field_names=("entity_id", "language"),
            fields=fields,
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)

        duplicate = run_composite(
            (
                {"item_id": "item-1", "entity_id": "card", "language": "hu"},
                {"item_id": "item-2", "entity_id": "card", "language": "hu"},
            ),
            field_names=("entity_id", "language"),
            fields=fields,
        )
        self.assertEqual(
            duplicate.diagnostics[0].observed_value,
            '[["string","card"],["language_tag","hu"]]',
        )


class TestCompositeUniquenessDuplicates(unittest.TestCase):
    def test_exact_two_and_three_record_groups_are_one_violation_each(self):
        for count in (2, 3):
            records = tuple(
                {"item_id": f"item-{index}", "group_id": "A", "sequence": 1}
                for index in range(count, 0, -1)
            )
            with self.subTest(count=count):
                result = run_composite(records)
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(result.evaluated_record_count, count)
                self.assertEqual(result.violation_count, 1)
                diagnostic = result.diagnostics[0]
                self.assertEqual(
                    diagnostic.code, execution.VALIDATION_UNIQUENESS_DUPLICATE
                )
                self.assertEqual(
                    diagnostic.observed_value,
                    '[["string","A"],["integer",1]]',
                )
                self.assertIn(
                    'registry:items.["group_id","sequence"]',
                    diagnostic.expected_contract,
                )
                self.assertEqual(
                    diagnostic.related_record_identities,
                    tuple(f"item-{index}" for index in range(1, count + 1)),
                )

    def test_two_duplicate_groups_are_reported_deterministically(self):
        result = run_composite(
            (
                {"item_id": "item-4", "group_id": "B", "sequence": 1},
                {"item_id": "item-2", "group_id": "A", "sequence": 1},
                {"item_id": "item-3", "group_id": "B", "sequence": 1},
                {"item_id": "item-1", "group_id": "A", "sequence": 1},
            )
        )

        self.assertEqual(result.violation_count, 2)
        self.assertEqual(
            tuple(item.related_record_identities for item in result.diagnostics),
            (("item-1", "item-2"), ("item-3", "item-4")),
        )

    def test_three_and_four_field_duplicates_are_detected(self):
        cases = (
            (
                ("ability_id", "parent_id", "sequence"),
                (
                    composite_field("ability_id"),
                    composite_field("parent_id"),
                    composite_field("sequence", "integer"),
                ),
                {"ability_id": "a1", "parent_id": "p1", "sequence": 1},
            ),
            (
                ("ability_id", "parent_id", "branch", "sequence"),
                (
                    composite_field("ability_id"),
                    composite_field("parent_id"),
                    composite_field("branch"),
                    composite_field("sequence", "integer"),
                ),
                {
                    "ability_id": "a1",
                    "parent_id": "p1",
                    "branch": "main",
                    "sequence": 1,
                },
            ),
        )
        for names, fields, values in cases:
            records = tuple(
                {"item_id": f"item-{index}", **values} for index in (1, 2)
            )
            with self.subTest(width=len(names)):
                result = run_composite(records, field_names=names, fields=fields)
                self.assertEqual(result.violation_count, 1)

    def test_lifecycle_fields_do_not_filter_duplicate_records(self):
        result = run_composite(
            (
                {
                    "item_id": "item-1", "group_id": "A", "sequence": 1,
                    "status": "active",
                },
                {
                    "item_id": "item-2", "group_id": "A", "sequence": 1,
                    "status": "archived", "export_enabled": False,
                },
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 2)
        self.assertEqual(result.violation_count, 1)


class TestCompositeUniquenessNullableD6(unittest.TestCase):
    def setUp(self):
        self.fields = (
            composite_field("ability_id"),
            composite_field("parent_id", nullable=True),
            composite_field("sequence", "integer"),
        )
        self.names = ("ability_id", "parent_id", "sequence")

    def test_explicit_null_is_a_real_duplicate_key_constituent(self):
        records = tuple(
            {
                "item_id": f"item-{index}",
                "ability_id": "a1",
                "parent_id": None,
                "sequence": 1,
            }
            for index in (2, 1)
        )

        result = run_composite(records, field_names=self.names, fields=self.fields)

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 2)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(
            result.diagnostics[0].observed_value,
            '[["string","a1"],["string",null],["integer",1]]',
        )

    def test_null_tuple_is_distinct_when_sequence_or_parent_changes(self):
        records = (
            {"item_id": "item-1", "ability_id": "a1", "parent_id": None, "sequence": 1},
            {"item_id": "item-2", "ability_id": "a1", "parent_id": None, "sequence": 2},
            {"item_id": "item-3", "ability_id": "a1", "parent_id": "p1", "sequence": 1},
        )

        result = run_composite(records, field_names=self.names, fields=self.fields)

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 3)

    def test_multiple_explicit_nulls_participate_in_grouping(self):
        names = ("ability_id", "parent_id", "branch", "sequence")
        fields = (
            composite_field("ability_id"),
            composite_field("parent_id", nullable=True),
            composite_field("branch", nullable=True),
            composite_field("sequence", "integer"),
        )
        records = (
            {"item_id": "item-1", "ability_id": "a1", "parent_id": None, "branch": None, "sequence": 1},
            {"item_id": "item-2", "ability_id": "a1", "parent_id": None, "branch": None, "sequence": 1},
            {"item_id": "item-3", "ability_id": "a1", "parent_id": None, "branch": "main", "sequence": 1},
            {"item_id": "item-4", "ability_id": "a1", "parent_id": None, "branch": None, "sequence": 2},
        )

        result = run_composite(records, field_names=names, fields=fields)

        self.assertEqual(result.violation_count, 1)
        self.assertEqual(result.evaluated_record_count, 4)
        self.assertIn('["string",null]', result.diagnostics[0].observed_value)

    def test_missing_field_is_malformed_and_never_merged_with_null(self):
        records = (
            {"item_id": "item-1", "ability_id": "a1", "parent_id": None, "sequence": 1},
            {"item_id": "item-2", "ability_id": "a1", "sequence": 1},
        )

        result = run_composite(records, field_names=self.names, fields=self.fields)

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 2)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(
            result.diagnostics[0].code,
            execution.VALIDATION_UNIQUENESS_TARGET_RECORD_INVALID,
        )
        self.assertEqual(
            result.diagnostics[0].reason, "target_field_missing:parent_id"
        )
        self.assertEqual(result.diagnostics[0].record_identity, "item-2")

    def test_null_in_non_nullable_schema_is_invalid(self):
        fields = (
            composite_field("ability_id"),
            composite_field("parent_id"),
            composite_field("sequence", "integer"),
        )
        result = run_composite(
            ({"item_id": "item-1", "ability_id": "a1", "parent_id": None, "sequence": 1},),
            field_names=self.names,
            fields=fields,
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(
            result.diagnostics[0].reason,
            "target_value_null_forbidden:parent_id",
        )


class TestNonNullCompositeUniquenessD7(unittest.TestCase):
    def test_one_field_complete_values_pass_and_duplicates_fail(self):
        passing = run_non_null(
            (
                {"item_id": "item-1", "value": "alpha"},
                {"item_id": "item-2", "value": "beta"},
            )
        )
        duplicate = run_non_null(
            (
                {"item_id": "item-1", "value": "alpha"},
                {"item_id": "item-2", "value": "alpha"},
            )
        )

        self.assertEqual(passing.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(passing.evaluated_record_count, 2)
        self.assertEqual(passing.diagnostics, ())
        self.assertEqual(duplicate.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(duplicate.evaluated_record_count, 2)
        self.assertEqual(duplicate.violation_count, 1)
        self.assertEqual(
            duplicate.diagnostics[0].code,
            execution.VALIDATION_UNIQUENESS_DUPLICATE,
        )

    def test_legitimate_nulls_are_evaluated_but_empty_applicable_set_passes(self):
        result = run_non_null(
            (
                {"item_id": "item-1", "value": None},
                {"item_id": "item-2", "value": None},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertIsNot(result.outcome, execution.ValidationOutcome.NOT_APPLICABLE)
        self.assertEqual(result.evaluated_record_count, 2)
        self.assertEqual(result.violation_count, 0)
        self.assertEqual(result.diagnostics, ())

    def test_partial_and_all_null_composite_keys_are_excluded(self):
        partial_fields = (
            composite_field("language", "language_tag"),
            composite_field("search_name", nullable=True),
        )
        partial = run_non_null(
            (
                {"item_id": "item-1", "language": "hu", "search_name": None},
                {"item_id": "item-2", "language": "hu", "search_name": None},
            ),
            field_names=("language", "search_name"),
            fields=partial_fields,
        )
        all_null = run_non_null(
            (
                {"item_id": "item-1", "left": None, "right": None},
                {"item_id": "item-2", "left": None, "right": None},
            ),
            field_names=("left", "right"),
        )

        for result in (partial, all_null):
            self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
            self.assertEqual(result.evaluated_record_count, 2)
            self.assertEqual(result.violation_count, 0)
            self.assertEqual(result.diagnostics, ())

    def test_d6_typed_null_participates_but_d7_null_is_excluded(self):
        names = ("group_id", "optional_id")
        fields = (
            composite_field("group_id"),
            composite_field("optional_id", nullable=True),
        )
        records = (
            {"item_id": "item-1", "group_id": "a", "optional_id": None},
            {"item_id": "item-2", "group_id": "a", "optional_id": None},
        )

        d6 = run_composite(records, field_names=names, fields=fields)
        d7 = run_non_null(records, field_names=names, fields=fields)

        self.assertEqual(d6.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(d6.violation_count, 1)
        self.assertEqual(d7.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(d7.violation_count, 0)
        self.assertEqual((d6.evaluated_record_count, d7.evaluated_record_count), (2, 2))

    def test_missing_is_malformed_and_never_treated_as_explicit_null(self):
        result = run_non_null(
            (
                {"item_id": "item-1", "value": None},
                {"item_id": "item-2"},
            )
        )

        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 2)
        self.assertEqual(result.violation_count, 1)
        self.assertEqual(
            result.diagnostics[0].code,
            execution.VALIDATION_UNIQUENESS_TARGET_RECORD_INVALID,
        )
        self.assertEqual(result.diagnostics[0].reason, "target_field_missing:value")
        self.assertEqual(result.diagnostics[0].record_identity, "item-2")

        invalid_identity = run_non_null(({"value": None},))
        self.assertEqual(invalid_identity.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(
            invalid_identity.diagnostics[0].reason,
            "target_identity_invalid",
        )

    def test_all_fields_are_validated_before_null_applicability(self):
        fields = (
            composite_field("optional_id", nullable=True),
            composite_field("sequence", "integer"),
        )
        for value in (True, 1.0, "1"):
            with self.subTest(value=value):
                result = run_non_null(
                    ({"item_id": "item-1", "optional_id": None, "sequence": value},),
                    field_names=("optional_id", "sequence"),
                    fields=fields,
                )
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(result.evaluated_record_count, 1)
                self.assertEqual(
                    result.diagnostics[0].reason,
                    "target_value_type_invalid:sequence",
                )

        missing_after_null = run_non_null(
            ({"item_id": "item-1", "optional_id": None},),
            field_names=("optional_id", "sequence"),
            fields=fields,
        )
        self.assertEqual(missing_after_null.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(
            missing_after_null.diagnostics[0].reason,
            "target_field_missing:sequence",
        )

        forbidden_null = run_non_null(
            ({"item_id": "item-1", "value": None},),
            fields=(composite_field("value"),),
        )
        self.assertEqual(forbidden_null.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(
            forbidden_null.diagnostics[0].reason,
            "target_value_null_forbidden:value",
        )

    def test_non_blocking_warning_duplicate_is_still_a_failure(self):
        rule = make_non_null_rule(
            blocking=False,
            severity_id="warning",
        )
        result = run_non_null(
            (
                {"item_id": "item-1", "value": "duplicate"},
                {"item_id": "item-2", "value": "duplicate"},
            ),
            rule=rule,
        )

        self.assertFalse(rule.blocking)
        self.assertEqual(rule.severity_id, "warning")
        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.violation_count, 1)

    def test_reversed_input_is_deterministic_with_excluded_and_duplicate_records(self):
        records = (
            {"item_id": "item-4", "value": None},
            {"item_id": "item-3", "value": "beta"},
            {"item_id": "item-2", "value": "alpha"},
            {"item_id": "item-1", "value": "alpha"},
        )

        forward = run_non_null(records)
        reverse = run_non_null(tuple(reversed(records)))

        self.assertEqual(forward, reverse)
        self.assertEqual(forward.evaluated_record_count, 4)
        self.assertEqual(forward.violation_count, 1)
        self.assertEqual(
            forward.diagnostics[0].related_record_identities,
            ("item-1", "item-2"),
        )

    def test_non_null_ast_contract_rejects_every_non_allowlisted_shape(self):
        literal = validation_expr.Literal("string", "value")
        field_list = validation_expr.ListLiteral((literal,))
        base = make_non_null_rule()
        invalid_items = (
            validation_expr.Call("trim", (literal,)),
            validation_expr.BinaryOperation("==", literal, literal),
            validation_expr.Binding("local", literal),
            validation_expr.MemberAccess(literal, "member"),
        )
        cases = [
            make_non_null_rule(()),
            make_non_null_rule(("value", "value")),
            make_uniqueness_rule(
                condition_expression='unique_non_null_by(["value",1])'
            ),
            replace(
                base,
                condition_ast=validation_expr.Call(
                    "unique_non_null_by", (field_list, field_list)
                ),
            ),
            replace(
                base,
                condition_ast=validation_expr.Call(
                    "unique_non_null_by", (literal,)
                ),
            ),
            replace(base, condition_ast=validation_expr.Sequence((base.condition_ast,))),
        ]
        cases.extend(
            replace(
                base,
                condition_ast=validation_expr.Call(
                    "unique_non_null_by",
                    (validation_expr.ListLiteral((item,)),),
                ),
            )
            for item in invalid_items
        )

        for rule in cases:
            with self.subTest(ast=rule.condition_ast):
                result = execution.execute_uniqueness_rule(
                    rule,
                    make_composite_context(
                        (), fields=(composite_field("value", nullable=True),)
                    ),
                )
                self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
                self.assertEqual(result.evaluated_record_count, 0)
                self.assertEqual(result.violation_count, 0)
                self.assertEqual(
                    result.diagnostics[0].code,
                    execution.VALIDATION_EXECUTOR_UNSUPPORTED,
                )

    def test_non_null_fields_must_resolve_once_to_supported_active_schema(self):
        unresolved = make_uniqueness_rule(
            condition_expression='unique_non_null_by(["value","missing"])'
        )
        ambiguous = make_non_null_rule(("group_id", "sequence"))
        unsupported = make_non_null_rule(("group_id", "sequence"))
        cases = (
            (
                unresolved,
                (composite_field("value", nullable=True),),
                "unique_non_null_by_field_unresolved:missing",
            ),
            (
                ambiguous,
                (
                    composite_field("group_id", field_id="fld_items_group_a"),
                    composite_field("group_id", field_id="fld_items_group_b"),
                    composite_field("sequence", "integer"),
                ),
                "unique_non_null_by_field_unresolved:group_id",
            ),
            (
                unsupported,
                (
                    composite_field("group_id", "boolean"),
                    composite_field("sequence", "integer"),
                ),
                "unique_non_null_by_field_schema_unsupported:group_id",
            ),
        )

        for rule, fields, reason in cases:
            with self.subTest(reason=reason):
                result = execution.execute_uniqueness_rule(
                    rule, make_composite_context((), fields=fields)
                )
                self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
                self.assertEqual(result.diagnostics[0].reason, reason)


class TestCompositeUniquenessTypes(unittest.TestCase):
    def test_materialized_types_are_exact_and_never_coerced(self):
        cases = (
            (
                {"item_id": "item-1", "group_id": "A", "sequence": True},
                None,
                "target_value_type_invalid:sequence",
            ),
            (
                {"item_id": "item-1", "group_id": "A", "sequence": 1.0},
                None,
                "target_value_type_invalid:sequence",
            ),
            (
                {"item_id": "item-1", "group_id": "A", "sequence": "1"},
                None,
                "target_value_type_invalid:sequence",
            ),
            (
                {"item_id": "item-1", "group_id": 1, "sequence": 1},
                None,
                "target_value_type_invalid:group_id",
            ),
            (
                {"item_id": "item-1", "entity_id": "card", "language": 1},
                (
                    composite_field("entity_id"),
                    composite_field("language", "language_tag"),
                ),
                "target_value_type_invalid:language",
            ),
        )
        for record, fields, reason in cases:
            names = ("entity_id", "language") if fields else ("group_id", "sequence")
            with self.subTest(reason=reason):
                result = run_composite(
                    (record,), field_names=names, fields=fields
                )
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(result.violation_count, 1)
                self.assertEqual(result.diagnostics[0].reason, reason)


class TestUniquenessDeterminism(unittest.TestCase):
    def test_reversed_input_has_identical_group_diagnostics(self):
        records = (
            {"item_id": "item-5", "value": "C"},
            {"item_id": "item-4", "value": "B"},
            {"item_id": "item-2", "value": "A"},
            {"item_id": "item-3", "value": "B"},
            {"item_id": "item-1", "value": "A"},
        )

        forward = run_uniqueness(records)
        reverse = run_uniqueness(tuple(reversed(records)))

        self.assertEqual(forward, reverse)
        self.assertEqual(
            tuple(item.observed_value for item in forward.diagnostics),
            ("A", "B"),
        )

    def test_reversed_malformed_records_have_identical_diagnostics(self):
        records = (
            {"value": "beta"},
            {"value": "alpha"},
        )

        forward = run_uniqueness(records)
        reverse = run_uniqueness(tuple(reversed(records)))

        self.assertEqual(forward, reverse)
        self.assertEqual(forward.violation_count, 2)

    def test_reversed_composite_input_has_identical_typed_diagnostics(self):
        records = (
            {"item_id": "item-4", "group_id": "B", "sequence": 1},
            {"item_id": "item-2", "group_id": "A", "sequence": 1},
            {"item_id": "item-3", "group_id": "B", "sequence": 1},
            {"item_id": "item-1", "group_id": "A", "sequence": 1},
        )

        forward = run_composite(records)
        reverse = run_composite(tuple(reversed(records)))

        self.assertEqual(forward, reverse)
        self.assertEqual(forward.violation_count, 2)


class TestDirectReferenceExecution(unittest.TestCase):
    def test_one_and_several_exact_references_pass(self):
        for records in (
            ({"item_id": "item-1", "reference_id": "alpha"},),
            (
                {"item_id": "item-2", "reference_id": "beta"},
                {"item_id": "item-1", "reference_id": "alpha"},
            ),
        ):
            with self.subTest(records=records):
                result = run_reference(
                    records,
                    reference_records=(
                        {"reference_id": "alpha", "status": "archived"},
                        {"reference_id": "beta", "status": "superseded"},
                    ),
                )
                self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
                self.assertEqual(result.evaluated_record_count, len(records))
                self.assertEqual(result.diagnostics, ())

    def test_missing_reference_fails(self):
        result = run_reference(
            ({"item_id": "item-1", "reference_id": "missing"},)
        )
        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.evaluated_record_count, 1)
        self.assertEqual(
            result.diagnostics[0].code, execution.VALIDATION_REFERENCE_MISSING
        )

    def test_null_missing_and_wrong_type_targets_fail(self):
        cases = (
            ({"item_id": "item-1", "reference_id": None}, "target_value_null"),
            ({"item_id": "item-1"}, "target_field_missing"),
            ({"item_id": "item-1", "reference_id": 1}, "target_value_invalid"),
        )
        for record, reason in cases:
            with self.subTest(reason=reason):
                result = run_reference((record,))
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(
                    result.diagnostics[0].code,
                    execution.VALIDATION_REFERENCE_TARGET_RECORD_INVALID,
                )
                self.assertEqual(result.diagnostics[0].reason, reason)

    def test_missing_target_primary_key_uses_content_identity_and_fails(self):
        result = run_reference(({"reference_id": "alpha", "other": 1},))
        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(
            result.diagnostics[0].reason, "target_primary_key_invalid"
        )
        self.assertTrue(
            result.diagnostics[0].record_identity.startswith("canonical-record:")
        )

    def test_missing_target_or_reference_infrastructure_fails_closed(self):
        cases = (
            (
                {"include_target": False},
                execution.VALIDATION_REFERENCE_TARGET_TABLE_MISSING,
            ),
            (
                {"include_reference": False},
                execution.VALIDATION_REFERENCE_TABLE_MISSING,
            ),
            (
                {"include_reference_schema": False},
                execution.VALIDATION_REFERENCE_FIELD_MISSING,
            ),
        )
        for options, code in cases:
            with self.subTest(code=code):
                result = run_reference((), **options)
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(result.diagnostics[0].code, code)

    def test_reference_record_missing_or_invalid_key_fails(self):
        for record in ({"status": "active"}, {"reference_id": None}, {"reference_id": 1}):
            with self.subTest(record=record):
                result = run_reference(
                    ({"item_id": "item-1", "reference_id": "alpha"},),
                    reference_records=(record,),
                )
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(
                    result.diagnostics[0].code,
                    execution.VALIDATION_REFERENCE_VALUE_INVALID,
                )

    def test_empty_target_is_vacuous_but_reference_contract_is_still_required(self):
        result = run_reference(())
        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 0)
        missing_reference = run_reference((), include_reference=False)
        self.assertEqual(missing_reference.outcome, execution.ValidationOutcome.FAIL)

    def test_empty_reference_fails_when_target_is_applicable(self):
        result = run_reference(
            ({"item_id": "item-1", "reference_id": "alpha"},),
            reference_records=(),
        )
        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.diagnostics[0].code, execution.VALIDATION_REFERENCE_MISSING)

    def test_duplicate_reference_is_ambiguous_not_first_or_last_wins(self):
        reference_records = (
            {"reference_id": "alpha", "status": "archived"},
            {"reference_id": "alpha", "status": "superseded"},
        )
        result = run_reference(
            ({"item_id": "item-1", "reference_id": "alpha"},),
            reference_records=reference_records,
        )
        reversed_result = run_reference(
            ({"item_id": "item-1", "reference_id": "alpha"},),
            reference_records=tuple(reversed(reference_records)),
        )
        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result, reversed_result)
        diagnostic = result.diagnostics[0]
        self.assertEqual(diagnostic.code, execution.VALIDATION_REFERENCE_AMBIGUOUS)
        self.assertEqual(diagnostic.related_record_identities, ("alpha", "alpha"))

    def test_case_whitespace_and_unicode_normalization_are_not_applied(self):
        cases = (
            ("Alpha", "alpha"),
            (" alpha ", "alpha"),
            ("e\u0301", "é"),
        )
        for observed, stored in cases:
            with self.subTest(observed=observed, stored=stored):
                result = run_reference(
                    ({"item_id": "item-1", "reference_id": observed},),
                    reference_records=({"reference_id": stored},),
                )
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(
                    result.diagnostics[-1].code,
                    execution.VALIDATION_REFERENCE_MISSING,
                )


SIMPLE_REFERENCE_GUARD = (
    'when(reference_id != null, '
    'exists("references","reference_id",reference_id))'
)
LEGACY_REFERENCE_GUARD = (
    'when(legacy_source_id not in [null,"#TBD"], '
    'exists("references","reference_id",legacy_source_id))'
)
DECISION_REFERENCE_GUARD = (
    'when(decision_id != null and starts_with(decision_id,"chg_"), '
    'exists("references","reference_id",decision_id))'
)


class TestGuardedReferenceExecution(unittest.TestCase):
    def test_simple_null_is_clause_satisfied_pass(self):
        result = run_reference(
            ({"item_id": "item-1", "reference_id": None},),
            expression=SIMPLE_REFERENCE_GUARD,
            reference_records=(),
        )
        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 0)
        self.assertIsNot(result.outcome, execution.ValidationOutcome.NOT_APPLICABLE)

    def test_simple_valid_and_dangling_non_null(self):
        passed = run_reference(
            ({"item_id": "item-1", "reference_id": "alpha"},),
            expression=SIMPLE_REFERENCE_GUARD,
        )
        failed = run_reference(
            ({"item_id": "item-1", "reference_id": "missing"},),
            expression=SIMPLE_REFERENCE_GUARD,
        )
        self.assertEqual(passed.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(passed.evaluated_record_count, 1)
        self.assertEqual(failed.outcome, execution.ValidationOutcome.FAIL)

    def test_legacy_null_and_tbd_pass_but_real_values_resolve(self):
        for value, outcome, evaluated in (
            (None, execution.ValidationOutcome.PASS, 0),
            ("#TBD", execution.ValidationOutcome.PASS, 0),
            ("src-1", execution.ValidationOutcome.PASS, 1),
            ("src-missing", execution.ValidationOutcome.FAIL, 1),
        ):
            with self.subTest(value=value):
                result = run_reference(
                    ({"item_id": "item-1", "legacy_source_id": value},),
                    expression=LEGACY_REFERENCE_GUARD,
                    target_field="legacy_source_id",
                    target_field_id="fld_items_legacy_source_id",
                    reference_records=({"reference_id": "src-1"},),
                )
                self.assertEqual(result.outcome, outcome)
                self.assertEqual(result.evaluated_record_count, evaluated)

    def test_decision_guard_has_exact_prefix_semantics(self):
        for value, outcome, evaluated in (
            (None, execution.ValidationOutcome.PASS, 0),
            ("manual_123", execution.ValidationOutcome.PASS, 0),
            ("chg_valid", execution.ValidationOutcome.PASS, 1),
            ("chg_missing", execution.ValidationOutcome.FAIL, 1),
        ):
            with self.subTest(value=value):
                result = run_reference(
                    ({"item_id": "item-1", "decision_id": value},),
                    expression=DECISION_REFERENCE_GUARD,
                    target_field="decision_id",
                    target_field_id="fld_items_decision_id",
                    reference_records=({"reference_id": "chg_valid"},),
                )
                self.assertEqual(result.outcome, outcome)
                self.assertEqual(result.evaluated_record_count, evaluated)

    def test_missing_guard_identifier_is_malformed_target_not_null(self):
        result = run_reference(
            ({"item_id": "item-1"},), expression=SIMPLE_REFERENCE_GUARD
        )
        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.diagnostics[0].reason, "target_field_missing")

    def test_non_allowlisted_or_mismatched_guard_shapes_are_unsupported(self):
        expressions = (
            'when(trim(reference_id) != null, exists("references","reference_id",reference_id))',
            'when(reference_id != null, exists("other","reference_id",reference_id))',
            'when(reference_id != null, exists("references","other_id",reference_id))',
            'when(reference_id != null, exists("references","reference_id",other_id))',
            'when(exists("references","reference_id",reference_id), exists("references","reference_id",reference_id))',
            'exists("references","reference_id",reference_id)',
            'when(reference_id == null, exists("references","reference_id",reference_id))',
            'when(starts_with(reference_id,"chg_"), exists("references","reference_id",reference_id))',
            'when(reference_id != "value", exists("references","reference_id",reference_id))',
            'when(reference_id not in [null,"OTHER"], exists("references","reference_id",reference_id))',
        )
        for expression in expressions:
            with self.subTest(expression=expression):
                rule = make_reference_rule(expression=expression)
                result = execution.execute_reference_integrity_rule(
                    rule, make_reference_context(guarded=True)
                )
                self.assertEqual(result.outcome, execution.ValidationOutcome.UNSUPPORTED)
                self.assertEqual(result.violation_count, 0)
                self.assertEqual(
                    result.diagnostics[0].code,
                    execution.VALIDATION_REFERENCE_GUARD_UNSUPPORTED,
                )


class TestReferenceDeterminism(unittest.TestCase):
    def test_target_and_reference_reordering_is_identical(self):
        targets = (
            {"item_id": "item-3", "reference_id": "missing-b"},
            {"item_id": "item-1", "reference_id": "alpha"},
            {"item_id": "item-2", "reference_id": "missing-a"},
        )
        references = (
            {"reference_id": "beta", "status": "superseded"},
            {"reference_id": "alpha", "status": "archived"},
        )
        forward = run_reference(targets, reference_records=references)
        reverse = run_reference(
            tuple(reversed(targets)), reference_records=tuple(reversed(references))
        )
        self.assertEqual(forward, reverse)
        self.assertEqual(
            tuple(item.record_identity for item in forward.diagnostics),
            ("item-2", "item-3"),
        )


def table_identity_namespace(
    namespace, target_values=(), *, include_target=True, include_schema_tables=True
):
    target_field_id = f"fld_{namespace}_items_target_table_id"
    table_field_id = f"fld_{namespace}_schema_tables_table_id"
    schema_tables = [
        {"table_id": "items", "primary_key": "item_id", "status": "active"},
        {
            "table_id": "schema_tables",
            "primary_key": "table_id",
            "status": "active",
        },
        {
            "table_id": "value_registry" if namespace == "registry" else "cards",
            "primary_key": "id",
            "status": "active",
        },
    ]
    tables = {
        "schema_fields": (
            {
                "field_id": target_field_id,
                "table_id": "items",
                "field_name": "target_table_id",
                "data_type": "string",
                "required_mode": "conditional",
                "nullable": True,
                "null_handling": "explicit_null",
                "status": "active",
            },
            {
                "field_id": table_field_id,
                "table_id": "schema_tables",
                "field_name": "table_id",
                "data_type": "string",
                "required_mode": "always",
                "nullable": False,
                "null_handling": "forbidden",
                "status": "active",
            },
        ),
    }
    if include_schema_tables:
        tables["schema_tables"] = tuple(schema_tables)
    if include_target:
        tables["items"] = tuple(
            {"item_id": f"{namespace}-{index}", "target_table_id": value}
            for index, value in enumerate(target_values, 1)
        )
    return tables


def table_identity_context(registry_values=(), carddatabase_values=(), **overrides):
    namespaced = {
        "registry": table_identity_namespace("registry", registry_values),
        "carddatabase": table_identity_namespace(
            "carddatabase", carddatabase_values
        ),
    }
    namespaced.update(overrides.pop("namespaced_overrides", {}))
    policies = {
        "registry": REFERENCE_POLICIES,
        "carddatabase": REFERENCE_POLICIES,
    }
    policies.update(overrides.pop("policy_overrides", {}))
    return execution.ValidationDataContext(
        namespaced_tables=namespaced,
        component_namespaces={
            "registry-component": "registry",
            "carddatabase-component": "carddatabase",
        },
        namespace_policies=policies,
        **overrides,
    )


def table_identity_rule(namespace="registry"):
    return make_reference_rule(
        expression=(
            'when(target_table_id != null, '
            'exists("schema_tables","table_id",target_table_id))'
        ),
        target_field="target_table_id",
        target_field_id=f"fld_{namespace}_items_target_table_id",
        reference_table="schema_tables",
        reference_field_id=f"fld_{namespace}_schema_tables_table_id",
        component=f"{namespace}-component",
    )


class TestReferenceNamespaces(unittest.TestCase):
    def test_same_table_id_coexists_and_unqualified_lookup_is_source_local(self):
        context = table_identity_context(
            registry_values=("value_registry",),
            carddatabase_values=("cards",),
        )
        registry = execution.execute_reference_integrity_rule(
            table_identity_rule("registry"), context
        )
        carddatabase = execution.execute_reference_integrity_rule(
            table_identity_rule("carddatabase"), context
        )
        self.assertEqual(registry.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(carddatabase.outcome, execution.ValidationOutcome.PASS)
        self.assertIsNone(context.table("schema_tables"))
        self.assertNotEqual(
            context.table("schema_tables", namespace="registry"),
            context.table("schema_tables", namespace="carddatabase"),
        )

    def test_qualified_table_identity_uses_target_namespace_not_raw_string(self):
        context = table_identity_context(
            registry_values=("carddatabase:cards",)
        )
        raw_registry_values = {
            row["table_id"]
            for row in context.table("schema_tables", namespace="registry")
        }
        self.assertNotIn("carddatabase:cards", raw_registry_values)
        result = execution.execute_reference_integrity_rule(
            table_identity_rule(), context
        )
        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(result.evaluated_record_count, 1)

    def test_unknown_malformed_or_uncontracted_namespace_fails_without_fallback(self):
        cases = (
            ("unknown:cards", {}, "qualified_namespace_unknown"),
            ("carddatabase:", {}, "qualified_identifier_malformed"),
            (
                "carddatabase:cards",
                {"policy_overrides": {"carddatabase": {}}},
                "qualified_table_identity_policy_missing",
            ),
        )
        for value, options, reason in cases:
            with self.subTest(value=value, reason=reason):
                result = execution.execute_reference_integrity_rule(
                    table_identity_rule(),
                    table_identity_context(registry_values=(value,), **options),
                )
                self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
                self.assertEqual(
                    result.diagnostics[-1].code,
                    execution.VALIDATION_REFERENCE_NAMESPACE_INVALID,
                )
                self.assertEqual(result.diagnostics[-1].reason, reason)

    def test_missing_qualified_table_identity_is_a_missing_reference(self):
        result = execution.execute_reference_integrity_rule(
            table_identity_rule(),
            table_identity_context(registry_values=("carddatabase:missing",)),
        )
        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(result.diagnostics[-1].code, execution.VALIDATION_REFERENCE_MISSING)

    def test_colon_is_not_a_global_namespace_heuristic(self):
        result = run_reference(
            ({"item_id": "item-1", "reference_id": "carddatabase:alpha"},),
            reference_records=({"reference_id": "carddatabase:alpha"},),
        )
        self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)

    def test_missing_source_binding_fails_closed(self):
        rule = make_reference_rule(component="unbound-component")
        result = execution.execute_reference_integrity_rule(
            rule, make_reference_context()
        )
        self.assertEqual(result.outcome, execution.ValidationOutcome.FAIL)
        self.assertEqual(
            result.diagnostics[0].code,
            execution.VALIDATION_REFERENCE_NAMESPACE_INVALID,
        )

    def test_duplicate_namespace_binding_is_rejected_at_context_boundary(self):
        with self.assertRaises(ValueError):
            execution.ValidationDataContext(
                namespaced_tables={"registry": reference_tables()},
                component_namespaces={"component-a": "registry", "component-b": "registry"},
                namespace_policies={"registry": REFERENCE_POLICIES},
            )

    def test_namespaced_context_requires_mapping_inputs(self):
        with self.assertRaises(ValueError):
            execution.ValidationDataContext(namespaced_tables=[])


def workbook_records(workbook, sheet_name, null_sentinel="#NULL"):
    rows = workbook[sheet_name].iter_rows(values_only=True)
    headers = tuple(next(rows))
    records = []
    for values in rows:
        if not any(value is not None for value in values):
            continue
        record = {}
        for field, value in zip(headers, values):
            if value == null_sentinel:
                value = None
            elif isinstance(value, float) and value.is_integer():
                value = int(value)
            elif isinstance(value, datetime):
                value = value.isoformat(timespec="seconds")
            elif isinstance(value, (date, time)):
                value = value.isoformat()
            record[field] = value
        records.append(record)
    return tuple(records)


def current_corpus_catalog_and_context(*, reverse=False):
    rule_inputs = []
    namespace_tables = []
    component_namespaces = []
    namespace_policies = []
    paths = tuple(reversed(WORKBOOKS)) if reverse else WORKBOOKS

    for path in paths:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            meta = {
                record["key"]: record["value"]
                for record in workbook_records(workbook, "META")
            }
            namespace = meta["export_namespace"]
            component_namespaces.append((path.name, namespace))
            namespace_policies.append(
                (
                    namespace,
                    {
                        "table_identity_policy": meta["table_identity_policy"],
                        "external_reference_identifier_policy": meta[
                            "external_reference_identifier_policy"
                        ],
                    },
                )
            )
            schemas = workbook_records(workbook, "SCHEMA_TABLES")
            tables = {
                schema["table_id"]: workbook_records(
                    workbook, schema["sheet_name"]
                )
                for schema in schemas
                if schema.get("status") == "active"
            }
            namespace_tables.append(
                (namespace, reverse_table_input(tables) if reverse else tables)
            )
            active_rules = tuple(
                record
                for record in workbook_records(workbook, "VALIDATION_RULES")
                if record.get("status") == "active"
            )
            if reverse:
                active_rules = tuple(reversed(active_rules))
            rule_inputs.extend(
                rules.ValidationRuleInput(record, path.name)
                for record in active_rules
            )
        finally:
            workbook.close()

    catalog_result = rules.build_validation_rule_catalog(tuple(rule_inputs))
    if not catalog_result.is_valid:
        raise AssertionError(catalog_result.diagnostics)
    return catalog_result.catalog, execution.ValidationDataContext(
        namespaced_tables=dict(namespace_tables),
        component_namespaces=dict(component_namespaces),
        namespace_policies=dict(namespace_policies),
    )


class TestCurrentWorkbookCorpus(unittest.TestCase):
    def test_current_corpus_discovers_and_executes_all_five_rules(self):
        hashes_before = {
            path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
        }
        rule_inputs = []
        registry_tables = {}
        allowed_counts = {}

        for path in WORKBOOKS:
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                raw_rules = workbook_records(workbook, "VALIDATION_RULES")
                active_rules = tuple(
                    record for record in raw_rules if record.get("status") == "active"
                )
                allowed_counts[path.name] = sum(
                    record.get("validation_kind_id") == "allowed_value"
                    for record in active_rules
                )
                rule_inputs.extend(
                    rules.ValidationRuleInput(record, path.name)
                    for record in active_rules
                )
                if path.name == "REGISTRY.xlsx":
                    for table_id, sheet_name in (
                        ("schema_tables", "SCHEMA_TABLES"),
                        ("schema_fields", "SCHEMA_FIELDS"),
                        ("value_groups", "VALUE_GROUPS"),
                        ("value_registry", "VALUE_REGISTRY"),
                        ("value_relations", "VALUE_RELATIONS"),
                        ("aliases", "ALIASES"),
                        ("localization", "LOCALIZATION"),
                    ):
                        registry_tables[table_id] = workbook_records(
                            workbook, sheet_name
                        )
            finally:
                workbook.close()

        catalog_result = rules.build_validation_rule_catalog(tuple(rule_inputs))
        self.assertTrue(catalog_result.is_valid, catalog_result.diagnostics)
        allowed_rules = catalog_result.catalog.by_kind("allowed_value")
        context = execution.ValidationDataContext(registry_tables)
        results = tuple(
            execution.execute_allowed_value_rule(rule, context)
            for rule in allowed_rules
        )

        self.assertEqual(allowed_counts, {"REGISTRY.xlsx": 5, "CARDDATABASE.xlsx": 0})
        self.assertEqual(len(allowed_rules), 5)
        self.assertEqual(len(results), 5)
        self.assertEqual(
            sum(result.outcome is execution.ValidationOutcome.UNSUPPORTED for result in results),
            0,
        )
        self.assertEqual(
            {result.rule_id: result.evaluated_record_count for result in results},
            {
                "val_aliases_alias_type_valid": 95,
                "val_aliases_normalization_mode_valid": 95,
                "val_localization_entity_type_valid": 167,
                "val_localization_language_valid": 167,
                "val_value_relations_relation_type_group": 59,
            },
        )
        for result in results:
            with self.subTest(rule_id=result.rule_id):
                self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
                self.assertEqual(result.violation_count, 0)
                self.assertEqual(result.diagnostics, ())

        hashes_after = {
            path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
        }
        self.assertEqual(hashes_before, hashes_after)


class TestCurrentRangeWorkbookCorpus(unittest.TestCase):
    def test_current_corpus_discovers_and_executes_all_ten_range_rules(self):
        hashes_before = {
            path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
        }
        rule_inputs = []
        context_tables = {"schema_tables": [], "schema_fields": []}
        range_counts = {}

        for path in WORKBOOKS:
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                schema_tables = workbook_records(workbook, "SCHEMA_TABLES")
                schema_fields = workbook_records(workbook, "SCHEMA_FIELDS")
                context_tables["schema_tables"].extend(schema_tables)
                context_tables["schema_fields"].extend(schema_fields)

                raw_rules = workbook_records(workbook, "VALIDATION_RULES")
                active_rules = tuple(
                    record for record in raw_rules if record.get("status") == "active"
                )
                range_rules = tuple(
                    record
                    for record in active_rules
                    if record.get("validation_kind_id") == "range"
                )
                range_counts[path.name] = len(range_rules)
                rule_inputs.extend(
                    rules.ValidationRuleInput(record, path.name)
                    for record in active_rules
                )

                schema_by_table = {
                    record["table_id"]: record
                    for record in schema_tables
                    if record.get("status") == "active"
                }
                for rule in range_rules:
                    table_id = rule["target_table_id"]
                    sheet_name = schema_by_table[table_id]["sheet_name"]
                    context_tables[table_id] = list(
                        workbook_records(workbook, sheet_name)
                    )
            finally:
                workbook.close()

        catalog_result = rules.build_validation_rule_catalog(tuple(rule_inputs))
        self.assertTrue(catalog_result.is_valid, catalog_result.diagnostics)
        range_rules = catalog_result.catalog.by_kind("range")
        context = execution.ValidationDataContext(context_tables)
        results = tuple(
            execution.execute_range_rule(rule, context) for rule in range_rules
        )

        self.assertEqual(range_counts, {"REGISTRY.xlsx": 1, "CARDDATABASE.xlsx": 9})
        self.assertEqual(len(range_rules), 10)
        self.assertEqual(len(results), 10)
        self.assertEqual(
            sum(result.outcome is execution.ValidationOutcome.UNSUPPORTED for result in results),
            0,
        )
        self.assertEqual(
            {result.rule_id: result.evaluated_record_count for result in results},
            {
                "cdb_val_card_keywords_sequence_positive": 399,
                "cdb_val_card_relations_sequence_positive": 0,
                "cdb_val_card_traits_sequence_positive": 0,
                "cdb_val_deck_entries_entry_index_positive": 32,
                "cdb_val_deck_entries_quantity_positive": 32,
                "cdb_val_decks_deck_index_positive": 2,
                "cdb_val_decks_required_card_count_positive": 2,
                "cdb_val_effect_parameter_item_index_positive": 14,
                "cdb_val_sets_release_order_positive": 3,
                "val_ability_template_nodes_output_sequence_positive": 22,
            },
        )
        for result in results:
            with self.subTest(rule_id=result.rule_id):
                self.assertEqual(result.outcome, execution.ValidationOutcome.PASS)
                self.assertEqual(result.violation_count, 0)
                self.assertEqual(result.diagnostics, ())

        hashes_after = {
            path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
        }
        self.assertEqual(hashes_before, hashes_after)


class TestCurrentReferenceWorkbookCorpus(unittest.TestCase):
    def test_current_corpus_executes_all_eleven_reference_rules(self):
        hashes_before = {
            path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
        }
        rule_inputs = []
        namespaced_tables = {}
        component_namespaces = {}
        policies = {}

        for path in WORKBOOKS:
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                meta = {
                    record["key"]: record["value"]
                    for record in workbook_records(workbook, "META")
                }
                namespace = meta["export_namespace"]
                component_namespaces[path.name] = namespace
                policies[namespace] = {
                    "table_identity_policy": meta["table_identity_policy"],
                    "external_reference_identifier_policy": meta[
                        "external_reference_identifier_policy"
                    ],
                }
                schema_tables = workbook_records(workbook, "SCHEMA_TABLES")
                tables = {}
                for schema in schema_tables:
                    if schema.get("status") == "active":
                        tables[schema["table_id"]] = workbook_records(
                            workbook, schema["sheet_name"]
                        )
                namespaced_tables[namespace] = tables
                active_rules = tuple(
                    record
                    for record in workbook_records(workbook, "VALIDATION_RULES")
                    if record.get("status") == "active"
                )
                rule_inputs.extend(
                    rules.ValidationRuleInput(record, path.name)
                    for record in active_rules
                )
            finally:
                workbook.close()

        catalog_result = rules.build_validation_rule_catalog(tuple(rule_inputs))
        self.assertTrue(catalog_result.is_valid, catalog_result.diagnostics)
        reference_rules = catalog_result.catalog.by_kind("reference_integrity")
        context = execution.ValidationDataContext(
            namespaced_tables=namespaced_tables,
            component_namespaces=component_namespaces,
            namespace_policies=policies,
        )
        results = tuple(
            execution.execute_reference_integrity_rule(rule, context)
            for rule in reference_rules
        )
        expected_counts = {
            "cdb_val_schema_fields_table_exists": 391,
            "val_contract_fields_schema_exists": 237,
            "val_export_manifest_table_exists": 30,
            "val_migration_map_changelog_decision_exists": 17,
            "val_migration_map_legacy_source_exists": 73,
            "val_migration_map_target_table_exists": 74,
            "val_restriction_types_blocked_action_exists": 0,
            "val_rule_registry_parent_exists": 13,
            "val_source_registry_supersedes_exists": 13,
            "val_value_relations_source_registry_value_exists": 59,
            "val_value_relations_target_registry_value_exists": 59,
        }
        target_count = 0
        for rule in reference_rules:
            namespace = context.namespace_for_component(rule.component_identity)
            target_count += len(context.table(rule.target_table_id, namespace=namespace))

        self.assertEqual(len(reference_rules), 11)
        self.assertEqual(sum(rule.condition_ast is None for rule in reference_rules), 5)
        self.assertEqual(sum(rule.condition_ast is not None for rule in reference_rules), 6)
        self.assertEqual(
            {result.rule_id: result.evaluated_record_count for result in results},
            expected_counts,
        )
        self.assertEqual(target_count, 1086)
        self.assertEqual(sum(expected_counts.values()), 966)
        self.assertEqual(target_count - sum(expected_counts.values()), 120)
        self.assertEqual(
            sum(result.outcome is execution.ValidationOutcome.UNSUPPORTED for result in results),
            0,
        )
        self.assertEqual(
            sum(result.outcome is execution.ValidationOutcome.FAIL for result in results),
            0,
        )
        self.assertTrue(all(result.outcome is execution.ValidationOutcome.PASS for result in results))
        self.assertEqual(sum(result.violation_count for result in results), 0)

        migration = next(
            rule
            for rule in reference_rules
            if rule.rule_id == "val_migration_map_target_table_exists"
        )
        migration_rows = context.table("migration_map", namespace="registry")
        local = sum(
            isinstance(row["target_table_id"], str)
            and ":" not in row["target_table_id"]
            for row in migration_rows
        )
        qualified = sum(
            isinstance(row["target_table_id"], str)
            and row["target_table_id"].startswith("carddatabase:")
            for row in migration_rows
        )
        guarded_out = sum(row["target_table_id"] is None for row in migration_rows)
        registry_raw = {
            row["table_id"]
            for row in context.table("schema_tables", namespace="registry")
        }
        self.assertEqual((local, qualified, guarded_out), (17, 57, 11))
        self.assertFalse(
            any(value.startswith("carddatabase:") for value in registry_raw)
        )
        migration_result = execution.execute_reference_integrity_rule(
            migration, context
        )
        self.assertEqual(migration_result.outcome, execution.ValidationOutcome.PASS)
        self.assertEqual(migration_result.evaluated_record_count, 74)

        hashes_after = {
            path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
        }
        self.assertEqual(hashes_before, hashes_after)


class TestCurrentUniquenessWorkbookCorpus(unittest.TestCase):
    def test_current_corpus_executes_all_73_uniqueness_rules(self):
        hashes_before = {
            path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
        }
        rule_inputs = []
        namespaced_tables = {}
        component_namespaces = {}

        for path in WORKBOOKS:
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                meta = {
                    record["key"]: record["value"]
                    for record in workbook_records(workbook, "META")
                }
                namespace = meta["export_namespace"]
                component_namespaces[path.name] = namespace
                schema_tables = workbook_records(workbook, "SCHEMA_TABLES")
                namespaced_tables[namespace] = {
                    schema["table_id"]: workbook_records(
                        workbook, schema["sheet_name"]
                    )
                    for schema in schema_tables
                    if schema.get("status") == "active"
                }
                active_rules = tuple(
                    record
                    for record in workbook_records(workbook, "VALIDATION_RULES")
                    if record.get("status") == "active"
                )
                rule_inputs.extend(
                    rules.ValidationRuleInput(record, path.name)
                    for record in active_rules
                )
            finally:
                workbook.close()

        catalog_result = rules.build_validation_rule_catalog(tuple(rule_inputs))
        self.assertTrue(catalog_result.is_valid, catalog_result.diagnostics)
        uniqueness_rules = catalog_result.catalog.by_kind("uniqueness")
        structured = tuple(
            rule
            for rule in uniqueness_rules
            if rule.condition_expression_source is None
        )
        expression_bearing = tuple(
            rule
            for rule in uniqueness_rules
            if rule.condition_expression_source is not None
        )
        composite = tuple(
            rule
            for rule in expression_bearing
            if rule.condition_expression_source.startswith("unique_by(")
        )
        unique_non_null = tuple(
            rule
            for rule in expression_bearing
            if rule.condition_expression_source.startswith("unique_non_null_by(")
        )
        context = execution.ValidationDataContext(
            namespaced_tables=namespaced_tables,
            component_namespaces=component_namespaces,
        )
        structured_results = tuple(
            execution.execute_uniqueness_rule(rule, context)
            for rule in structured
        )
        composite_results = tuple(
            execution.execute_uniqueness_rule(rule, context)
            for rule in composite
        )
        non_null_results = tuple(
            execution.execute_uniqueness_rule(rule, context)
            for rule in unique_non_null
        )

        self.assertEqual(len(catalog_result.catalog), 308)
        self.assertEqual(
            sum(rule.blocking for rule in catalog_result.catalog.rules),
            306,
        )
        self.assertEqual(len(uniqueness_rules), 73)
        self.assertEqual(len(structured), 39)
        self.assertEqual(len(expression_bearing), 34)
        self.assertEqual(len(composite), 28)
        self.assertEqual(len(unique_non_null), 6)
        self.assertTrue(
            all(
                result.outcome is execution.ValidationOutcome.PASS
                for result in structured_results + composite_results
            )
        )
        self.assertTrue(
            all(
                result.diagnostics == ()
                for result in structured_results + composite_results
            )
        )
        self.assertEqual(
            sum(result.evaluated_record_count for result in structured_results),
            4274,
        )
        self.assertEqual(
            sum(result.evaluated_record_count for result in composite_results),
            3413,
        )
        self.assertEqual(
            sum(
                result.violation_count
                for result in structured_results + composite_results
            ),
            0,
        )

        null_count = 0
        missing_count = 0
        primary_key_rules = []
        non_primary_rules = []
        empty_tables = set()
        for rule, result in zip(structured, structured_results):
            namespace = context.namespace_for_component(rule.component_identity)
            field_schema = core._resolve_field_schema(
                context,
                rule.target_table_id,
                field_id=rule.target_field_id,
                namespace=namespace,
            )
            primary_key = core._resolve_primary_key(
                context, rule.target_table_id, namespace=namespace
            )
            target_records = core._resolve_table(
                context, rule.target_table_id, namespace=namespace
            )
            self.assertEqual(
                (
                    field_schema["data_type"],
                    field_schema["required_mode"],
                    field_schema["nullable"],
                    field_schema["null_handling"],
                ),
                ("string", "always", False, "forbidden"),
            )
            field_name = field_schema["field_name"]
            missing_count += sum(field_name not in record for record in target_records)
            null_count += sum(
                field_name in record and record[field_name] is None
                for record in target_records
            )
            if field_name == primary_key:
                primary_key_rules.append(rule.rule_id)
            else:
                non_primary_rules.append(rule.rule_id)
            if result.evaluated_record_count == 0:
                empty_tables.add(rule.target_table_id)

        self.assertEqual(null_count, 0)
        self.assertEqual(missing_count, 0)
        self.assertEqual(len(primary_key_rules), 37)
        self.assertEqual(
            sorted(non_primary_rules),
            [
                "cdb_val_schema_tables_sheet_name_unique",
                "cdb_val_sets_set_code_unique",
            ],
        )
        result_by_id = {result.rule_id: result for result in structured_results}
        for rule_id in non_primary_rules:
            self.assertEqual(
                result_by_id[rule_id].outcome, execution.ValidationOutcome.PASS
            )
        self.assertEqual(
            empty_tables,
            {
                "card_relations",
                "card_traits",
                "ability_choice_options",
                "ability_choices",
                "ability_costs",
                "effect_tags",
                "ability_usage_limits",
            },
        )
        self.assertTrue(
            all(
                result.evaluated_record_count == 0
                and result.violation_count == 0
                and result.outcome is execution.ValidationOutcome.PASS
                for rule, result in zip(structured, structured_results)
                if rule.target_table_id in empty_tables
            )
        )

        nullable_stats = {}
        one_null_count = 0
        multiple_null_count = 0
        for rule, result in zip(composite, composite_results):
            namespace = context.namespace_for_component(rule.component_identity)
            field_names = tuple(
                item.value for item in rule.condition_ast.arguments[0].items
            )
            field_schemas = tuple(
                core._resolve_field_schema(
                    context,
                    rule.target_table_id,
                    field_name=field_name,
                    namespace=namespace,
                )
                for field_name in field_names
            )
            target_records = core._resolve_table(
                context, rule.target_table_id, namespace=namespace
            )
            self.assertTrue(all(schema is not None for schema in field_schemas))
            self.assertTrue(
                all(field_name in record for record in target_records for field_name in field_names)
            )
            null_widths = tuple(
                sum(record[field_name] is None for field_name in field_names)
                for record in target_records
            )
            one_null = sum(width == 1 for width in null_widths)
            multiple_null = sum(width > 1 for width in null_widths)
            one_null_count += one_null
            multiple_null_count += multiple_null
            if any(schema["nullable"] is True for schema in field_schemas):
                nullable_stats[rule.rule_id] = (
                    len(target_records),
                    one_null + multiple_null,
                    result.violation_count,
                    result.outcome,
                )

        self.assertEqual(one_null_count, 12)
        self.assertEqual(multiple_null_count, 21)
        self.assertEqual(
            nullable_stats,
            {
                "cdb_val_conditions_sequence_unique": (
                    4, 4, 0, execution.ValidationOutcome.PASS,
                ),
                "cdb_val_effects_sequence_unique": (
                    21, 21, 0, execution.ValidationOutcome.PASS,
                ),
                "cdb_val_expressions_sequence_unique": (
                    9, 8, 0, execution.ValidationOutcome.PASS,
                ),
                "cdb_val_effect_tags_tuple_unique": (
                    0, 0, 0, execution.ValidationOutcome.PASS,
                ),
            },
        )

        expected_non_null_outcomes = {
            "cdb_val_card_localization_search_name_collision": (
                execution.ValidationOutcome.FAIL, 814, 1,
            ),
            "cdb_val_export_manifest_file_unique": (
                execution.ValidationOutcome.PASS, 30, 0,
            ),
            "cdb_val_export_manifest_order_unique": (
                execution.ValidationOutcome.PASS, 30, 0,
            ),
            "val_export_manifest_load_order_unique": (
                execution.ValidationOutcome.PASS, 30, 0,
            ),
            "val_export_manifest_output_name_unique": (
                execution.ValidationOutcome.PASS, 30, 0,
            ),
            "val_rule_registry_engine_module_unique": (
                execution.ValidationOutcome.PASS, 23, 0,
            ),
        }
        self.assertEqual(
            {
                result.rule_id: (
                    result.outcome,
                    result.evaluated_record_count,
                    result.violation_count,
                )
                for result in non_null_results
            },
            expected_non_null_outcomes,
        )

        excluded_null = 0
        grouped_non_null = 0
        missing_constituents = 0
        for rule in unique_non_null:
            namespace = context.namespace_for_component(rule.component_identity)
            field_names = tuple(
                item.value for item in rule.condition_ast.arguments[0].items
            )
            target_records = core._resolve_table(
                context, rule.target_table_id, namespace=namespace
            )
            for record in target_records:
                missing = any(field_name not in record for field_name in field_names)
                missing_constituents += missing
                if missing:
                    continue
                contains_null = any(record[field_name] is None for field_name in field_names)
                excluded_null += contains_null
                grouped_non_null += not contains_null

        self.assertEqual(
            sum(result.evaluated_record_count for result in non_null_results),
            957,
        )
        self.assertEqual((grouped_non_null, excluded_null, missing_constituents), (951, 6, 0))

        all_uniqueness_results = (
            structured_results + composite_results + non_null_results
        )
        self.assertEqual(
            sum(
                result.outcome is execution.ValidationOutcome.PASS
                for result in all_uniqueness_results
            ),
            72,
        )
        self.assertEqual(
            sum(
                result.outcome is execution.ValidationOutcome.FAIL
                for result in all_uniqueness_results
            ),
            1,
        )
        self.assertEqual(
            sum(
                result.outcome is execution.ValidationOutcome.UNSUPPORTED
                for result in all_uniqueness_results
            ),
            0,
        )
        collision_rule = next(
            rule
            for rule in unique_non_null
            if rule.rule_id == "cdb_val_card_localization_search_name_collision"
        )
        collision_result = next(
            result
            for result in non_null_results
            if result.rule_id == collision_rule.rule_id
        )
        self.assertFalse(collision_rule.blocking)
        self.assertEqual(collision_rule.severity_id, "warning")
        self.assertEqual(
            collision_result.diagnostics[0].observed_value,
            '[["language_tag","hu"],["string","viharmadár"]]',
        )
        self.assertEqual(
            collision_result.diagnostics[0].related_record_identities,
            ("cardloc_ter_whu_021_hu", "cardloc_ven_egu_007_hu"),
        )

        executable_counts = (
            len(catalog_result.catalog.by_kind("allowed_value")),
            len(catalog_result.catalog.by_kind("range")),
            len(catalog_result.catalog.by_kind("reference_integrity")),
            len(structured),
            len(composite),
            len(unique_non_null),
            len(catalog_result.catalog.by_kind("hierarchy_integrity")),
            len(catalog_result.catalog.by_kind("contract_field_invariant")),
            len(catalog_result.catalog.by_kind("qualified_reference_match")),
        )
        self.assertEqual(executable_counts, (5, 10, 11, 39, 28, 6, 7, 3, 7))
        self.assertEqual(sum(executable_counts), 116)
        executable_rules = tuple(
            rule
            for rule in catalog_result.catalog.rules
            if rule.validation_kind_id
            in {
                "allowed_value",
                "contract_field_invariant",
                "hierarchy_integrity",
                "qualified_reference_match",
                "range",
                "reference_integrity",
                "uniqueness",
            }
        )
        self.assertEqual(len(executable_rules), 116)
        self.assertEqual(sum(rule.blocking for rule in executable_rules), 115)

        hashes_after = {
            path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
        }
        self.assertEqual(hashes_before, hashes_after)


class TestUnifiedCurrentWorkbookCorpus(unittest.TestCase):
    @staticmethod
    def execute_current_rules(catalog, context):
        executors = {
            "allowed_value": execution.execute_allowed_value_rule,
            "contract_field_invariant": execution.execute_contract_field_invariant_rule,
            "hierarchy_integrity": execution.execute_hierarchy_integrity_rule,
            "qualified_reference_match": execution.execute_qualified_reference_match_rule,
            "range": execution.execute_range_rule,
            "reference_integrity": execution.execute_reference_integrity_rule,
            "uniqueness": execution.execute_uniqueness_rule,
        }
        selected = tuple(
            rule
            for rule in catalog.by_stage("pre_export")
            if rule.validation_kind_id in executors
        )
        return selected, tuple(
            executors[rule.validation_kind_id](rule, context)
            for rule in selected
        )

    def test_all_116_executable_rules_share_one_unified_context(self):
        hashes_before = {
            path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
        }
        catalog, context = current_corpus_catalog_and_context()
        selected, results = self.execute_current_rules(catalog, context)
        family_results = {
            kind: tuple(
                result
                for rule, result in zip(selected, results)
                if rule.validation_kind_id == kind
            )
            for kind in (
                "allowed_value",
                "contract_field_invariant",
                "hierarchy_integrity",
                "qualified_reference_match",
                "range",
                "reference_integrity",
                "uniqueness",
            )
        }

        self.assertEqual(len(selected), 116)
        self.assertEqual(
            Counter(result.outcome for result in results),
            {
                execution.ValidationOutcome.PASS: 115,
                execution.ValidationOutcome.FAIL: 1,
            },
        )
        self.assertEqual(
            {
                kind: Counter(result.outcome for result in values)
                for kind, values in family_results.items()
            },
            {
                "allowed_value": {execution.ValidationOutcome.PASS: 5},
                "contract_field_invariant": {
                    execution.ValidationOutcome.PASS: 3
                },
                "hierarchy_integrity": {execution.ValidationOutcome.PASS: 7},
                "qualified_reference_match": {
                    execution.ValidationOutcome.PASS: 7
                },
                "range": {execution.ValidationOutcome.PASS: 10},
                "reference_integrity": {execution.ValidationOutcome.PASS: 11},
                "uniqueness": {
                    execution.ValidationOutcome.PASS: 72,
                    execution.ValidationOutcome.FAIL: 1,
                },
            },
        )
        self.assertFalse(
            any(
                diagnostic.reason == "target_field_unresolved"
                for kind in ("allowed_value", "range")
                for result in family_results[kind]
                for diagnostic in result.diagnostics
            )
        )
        failures = tuple(
            (rule, result)
            for rule, result in zip(selected, results)
            if result.outcome is execution.ValidationOutcome.FAIL
        )
        self.assertEqual(len(failures), 1)
        failed_rule, failed_result = failures[0]
        self.assertEqual(
            failed_rule.rule_id,
            "cdb_val_card_localization_search_name_collision",
        )
        self.assertFalse(failed_rule.blocking)
        self.assertEqual(failed_rule.severity_id, "warning")
        self.assertEqual(failed_result.violation_count, 1)
        self.assertEqual(
            sum(
                result.outcome is execution.ValidationOutcome.UNSUPPORTED
                for result in results
            ),
            0,
        )
        hashes_after = {
            path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS
        }
        self.assertEqual(hashes_before, hashes_after)

    def test_reversed_workbook_and_namespace_order_is_deterministic(self):
        forward_catalog, forward_context = current_corpus_catalog_and_context()
        reverse_catalog, reverse_context = current_corpus_catalog_and_context(
            reverse=True
        )

        forward = self.execute_current_rules(forward_catalog, forward_context)
        reverse = self.execute_current_rules(reverse_catalog, reverse_context)

        self.assertEqual(
            tuple(rule.rule_id for rule in forward[0]),
            tuple(rule.rule_id for rule in reverse[0]),
        )
        self.assertEqual(forward[1], reverse[1])


if __name__ == "__main__":
    unittest.main()
