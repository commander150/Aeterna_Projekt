import hashlib
import importlib.util
import re
import sys
import unittest
from dataclasses import FrozenInstanceError
from pathlib import Path

from openpyxl import load_workbook


PYTHON_ROOT = Path(__file__).resolve().parents[1]
REPOSITORY_ROOT = Path(__file__).resolve().parents[3]
MODULE_DIRECTORY = PYTHON_ROOT / "tools" / "canonical_export"
SCRIPT_PATH = MODULE_DIRECTORY / "canonical_validation_expr.py"
WORKBOOKS = (
    REPOSITORY_ROOT / "data" / "canonical" / "REGISTRY.xlsx",
    REPOSITORY_ROOT / "data" / "canonical" / "CARDDATABASE.xlsx",
)


def load_validation_module():
    module_name = "canonical_validation_expr"
    sys.path.insert(0, str(MODULE_DIRECTORY))
    try:
        spec = importlib.util.spec_from_file_location(module_name, SCRIPT_PATH)
        module = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = module
        spec.loader.exec_module(module)
        return module
    finally:
        sys.path.remove(str(MODULE_DIRECTORY))


validation = load_validation_module()


def error_code(source, *, limits=validation.DEFAULT_LIMITS):
    with unittest.TestCase().assertRaises(validation.ValidationExpressionError) as raised:
        validation.parse_validation_expression(source, limits=limits)
    return raised.exception.diagnostics[0].code


def active_validation_corpus():
    active_count = 0
    structured_only_count = 0
    expressions = []
    hashes_before = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS}
    for path in WORKBOOKS:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = workbook["VALIDATION_RULES"].iter_rows(values_only=True)
            headers = next(rows)
            index = {name: position for position, name in enumerate(headers)}
            for row_number, row in enumerate(rows, start=2):
                if row[index["status"]] != "active":
                    continue
                active_count += 1
                expression = row[index["condition_expression"]]
                if isinstance(expression, str) and expression not in {"", "#NULL"}:
                    expressions.append(
                        (
                            path.name,
                            row_number,
                            row[index["validation_rule_id"]],
                            expression,
                        )
                    )
                else:
                    structured_only_count += 1
        finally:
            workbook.close()
    hashes_after = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in WORKBOOKS}
    if hashes_before != hashes_after:
        raise AssertionError("Read-only corpus loading changed a workbook.")
    return active_count, structured_only_count, tuple(expressions)


class TestTokenizerAndGrammar(unittest.TestCase):
    def test_longest_operators_keywords_and_source_positions(self):
        tokens = validation.tokenize_validation_expression(
            'a != b\nand c <= 2 or c >= 1 and c == 2'
        )
        self.assertEqual(
            [token.kind for token in tokens],
            [
                "IDENTIFIER",
                "NE",
                "IDENTIFIER",
                "AND",
                "IDENTIFIER",
                "LTE",
                "INTEGER",
                "OR",
                "IDENTIFIER",
                "GTE",
                "INTEGER",
                "AND",
                "IDENTIFIER",
                "EQ",
                "INTEGER",
                "EOF",
            ],
        )
        self.assertEqual((tokens[3].line, tokens[3].column), (2, 1))

    def test_explicit_operator_precedence(self):
        root = validation.parse_validation_expression(
            "not a + 1 == b and c or d", validate_builtins=False
        )
        projection = validation.canonical_ast_projection(root)["root"]
        self.assertEqual(projection["operator"], "or")
        self.assertEqual(projection["left"]["operator"], "and")
        comparison = projection["left"]["left"]
        self.assertEqual(comparison["operator"], "==")
        self.assertEqual(comparison["left"]["operator"], "+")
        self.assertEqual(comparison["left"]["left"]["operator"], "not")

    def test_membership_addition_and_member_access(self):
        root = validation.parse_validation_expression(
            'lookup_record("items","id",item_id).status not in ["old","deleted"]'
        )
        projection = validation.canonical_ast_projection(root)["root"]
        self.assertEqual(projection["operator"], "not in")
        self.assertEqual(projection["left"]["node"], "member")
        self.assertEqual(projection["left"]["member"], "status")

    def test_signed_int64_domain(self):
        minimum = validation.parse_validation_expression(
            "-9223372036854775808", validate_builtins=False
        )
        maximum = validation.parse_validation_expression(
            "+9223372036854775807", validate_builtins=False
        )
        self.assertEqual(minimum.value, -(2**63))
        self.assertEqual(maximum.value, 2**63 - 1)
        self.assertEqual(
            error_code("9223372036854775808"), validation.INTEGER_RANGE
        )
        self.assertEqual(
            error_code("-9223372036854775809"), validation.INTEGER_RANGE
        )

    def test_json_strings_and_unicode_scalar_contract(self):
        escaped = validation.parse_validation_expression(
            r'label == "\u0150rz\u0151 \uD83D\uDE00"',
            validate_builtins=False,
        )
        literal = escaped.right
        self.assertEqual(literal.value, "\u0150rz\u0151 \U0001f600")
        self.assertEqual(
            error_code(r'label == "\uD800"'), validation.LITERAL_INVALID
        )
        self.assertEqual(
            error_code('label == "line\nbreak"'), validation.LITERAL_INVALID
        )

    def test_tbd_is_not_a_string_or_identifier(self):
        bare = validation.parse_validation_expression(
            "x == tbd", validate_builtins=False
        )
        sentinel = validation.parse_validation_expression(
            'x == "#TBD"', validate_builtins=False
        )
        self.assertIsInstance(bare.right, validation.TbdLiteral)
        self.assertIsInstance(sentinel.right, validation.Literal)
        self.assertNotEqual(
            validation.semantic_ast_hash(bare),
            validation.semantic_ast_hash(sentinel),
        )

    def test_binding_and_sequence_are_distinct_from_comparison(self):
        root = validation.parse_validation_expression(
            'x = lookup_record("items","id",item_id); x != null'
        )
        self.assertIsInstance(root, validation.Sequence)
        self.assertIsInstance(root.statements[0], validation.Binding)
        self.assertIsInstance(root.statements[1], validation.BinaryOperation)
        equality = validation.parse_validation_expression(
            "x == 1", validate_builtins=False
        )
        self.assertIsInstance(equality, validation.BinaryOperation)

    def test_when_constraint_body_can_contain_bindings(self):
        root = validation.parse_validation_expression(
            'when(enabled == true, row = lookup_record("items","id",item_id); '
            'row != null and row.status == "active")'
        )
        self.assertIsInstance(root, validation.Call)
        self.assertEqual(root.function, "when")
        self.assertIsInstance(root.arguments[1], validation.Sequence)
        self.assertIsInstance(root.arguments[1].statements[0], validation.Binding)

    def test_keyword_argument_and_member_assignment_are_rejected(self):
        self.assertEqual(
            error_code('lookup(table = "items")'), validation.UNSUPPORTED_SYNTAX
        )
        self.assertEqual(error_code("record.field = 1"), validation.UNSUPPORTED_SYNTAX)


class TestImmutableAstAndMap(unittest.TestCase):
    def test_ast_nodes_and_catalog_are_immutable(self):
        identifier = validation.Identifier("value")
        with self.assertRaises(FrozenInstanceError):
            identifier.name = "changed"
        with self.assertRaises(FrozenInstanceError):
            validation.BUILTIN_CATALOG.entries = ()
        list_node = validation.ListLiteral((identifier,))
        self.assertIsInstance(list_node.items, tuple)

    def test_restricted_map_is_string_to_string_and_byte_sorted(self):
        first = validation.parse_validation_expression(
            r'map(k,{"\uD800\uDC00":"supplementary","\uE000":"private"})'
        )
        second = validation.parse_validation_expression(
            r'map(k,{"\uE000":"private","\uD800\uDC00":"supplementary"})'
        )
        map_node = first.arguments[1]
        self.assertIsInstance(map_node, validation.MapLiteral)
        self.assertEqual(map_node.items, (("\ue000", "private"), ("\U00010000", "supplementary")))
        self.assertEqual(
            validation.semantic_ast_hash(first), validation.semantic_ast_hash(second)
        )
        self.assertEqual(error_code('{"a":1}'), validation.SYNTAX_ERROR)
        self.assertEqual(error_code('{"a":"1","a":"2"}'), validation.LITERAL_INVALID)

    def test_current_corpus_map_shape_is_exact(self):
        _, _, corpus = active_validation_corpus()
        source = next(
            expression
            for _, _, rule_id, expression in corpus
            if rule_id == "val_value_relations_target_group_mapping"
        )
        root = validation.parse_validation_expression(source)
        map_node = root.statements[0].value.arguments[1]
        expected = {
            "accepts_damage_kind": "damage_kind",
            "accepts_play_entry_method": "play_entry_method",
            "continues_toward_zone": "zone",
            "persists_while_in_zone": "zone",
            "requires_actual_destination_zone": "zone",
            "requires_card_relation": "card_relation_type",
            "requires_damage_kind": "damage_kind",
            "requires_matching_effective_keyword": "keyword",
            "requires_phase": "phase",
            "requires_resulting_destruction_cause_kind": "destruction_cause_kind",
            "requires_source_domain_row": "domain_row",
            "requires_subject_activity_state": "activity_state",
            "requires_subject_domain_row": "domain_row",
            "requires_target_domain_row": "domain_row",
            "uses_destruction_cause_kind": "destruction_cause_kind",
            "uses_obligation_requirement": "obligation_requirement",
            "uses_stacking_policy": "stacking_policy",
        }
        self.assertEqual(dict(map_node.items), expected)
        self.assertEqual(
            list(map_node.items),
            sorted(map_node.items, key=lambda item: item[0].encode("utf-8")),
        )


class TestBuiltinCatalog(unittest.TestCase):
    def test_catalog_has_exact_current_name_set_and_metadata(self):
        self.assertEqual(len(validation.BUILTIN_CATALOG), 54)
        for entry in validation.BUILTIN_CATALOG.entries:
            self.assertTrue(entry.name)
            self.assertTrue(entry.allowed_arities)
            self.assertEqual(tuple(sorted(set(entry.allowed_arities))), entry.allowed_arities)
            self.assertIn(
                entry.result_category,
                {
                    "boolean",
                    "integer",
                    "string",
                    "record",
                    "scalar",
                    "collection",
                    "version",
                    "unknown/deferred",
                },
            )
            self.assertIn(entry.execution_domain, {"static", "runtime", "mixed", "structural"})
            self.assertIn(
                entry.semantic_status,
                {"confirmed", "strongly_inferred", "execution_deferred"},
            )

    def test_observed_overload_and_arity_validation(self):
        validation.parse_validation_expression("count(values)")
        validation.parse_validation_expression('count("items", status == "active")')
        self.assertEqual(error_code("count()"), validation.ARITY_MISMATCH)
        self.assertEqual(error_code("count(a,b,c)"), validation.ARITY_MISMATCH)

    def test_unknown_builtin_is_structured_and_does_not_execute(self):
        with self.assertRaises(validation.ValidationExpressionError) as raised:
            validation.parse_validation_expression("not_a_builtin(value)")
        diagnostic = raised.exception.diagnostics[0]
        self.assertEqual(diagnostic.code, validation.UNKNOWN_BUILTIN)
        self.assertEqual(diagnostic.as_dict()["context"]["builtin"], "not_a_builtin")

    def test_when_contract_is_metadata_not_evaluation(self):
        signature = validation.BUILTIN_CATALOG.get("when")
        self.assertEqual(signature.semantic_status, "confirmed")
        self.assertEqual(validation.WHEN_FALSE_GUARD_OUTCOME, "constraint_satisfied")
        self.assertFalse(validation.WHEN_PRODUCES_NOT_APPLICABLE)
        self.assertEqual(validation.SEQUENCE_CONSTRAINT_COMBINATION, "conjunction")
        root = validation.parse_validation_expression("when(a == 1, b == 2); c == 3")
        self.assertIsInstance(root, validation.Sequence)


class TestCanonicalAstIdentity(unittest.TestCase):
    GOLDEN_HASHES = {
        "a == b": "sha256:d5fdf3a325d426c231150aad44e876ef7cd6b16907c6cf626d8059d3634dcd61",
        "when(a == 1, b == 2)": "sha256:46f163ca820fbcfe685f56e47f2f4e1b4c7c27a654586170e3aa87fc3e894752",
        'x = lookup_record("items","id",item_id); x != null': "sha256:b26017a188113fd9053f8095264c610fd9e39cf110f325b7ff885f462006fffc",
        'lookup_record("items","id",item_id).status == "active"': "sha256:89985834fc76943952c3b68103fb363fdd4d2bf84fc602a5c3c74ca27650d793",
        'unique_by(["card_id","ability_index"])': "sha256:7ad10b6bb5f8ed2140a9f0309759bd2af56f8b17a1a7c6c6a461ac16714284b1",
        'hierarchy_is_acyclic("rule_id","parent_rule_id") and no_self_reference("rule_id","parent_rule_id")': "sha256:a0a51d4f7db731508f5f81aa041ad171ae6b6f541b794ff61edd68c243beb268",
        'applied = prior_event("event_modifier_applied","modifier_instance_id",modifier_instance_id); applied != null': "sha256:e9d342b187c16b04556e1fb595095c5a92f60ba3327fa565bcb352baf2646af1",
        r'label == "\u0150rz\u0151 \uD83D\uDE00"': "sha256:f30d2ca20d5808f2909bcaaca5281ad5cbb307f40259ccdcf036450f9bdc9a3a",
    }

    def test_projection_is_explicit_and_language_neutral(self):
        root = validation.parse_validation_expression("a == b", validate_builtins=False)
        self.assertEqual(
            validation.canonical_ast_projection(root),
            {
                "ast_format": "aeterna_validation_ast_v1",
                "expression_language": "aeterna_validation_expr_v1",
                "root": {
                    "node": "binary",
                    "operator": "==",
                    "left": {"node": "identifier", "name": "a"},
                    "right": {"node": "identifier", "name": "b"},
                },
            },
        )

    def test_representative_golden_hash_vectors(self):
        for source, expected in self.GOLDEN_HASHES.items():
            with self.subTest(source=source):
                root = validation.parse_validation_expression(source)
                self.assertEqual(validation.semantic_ast_hash(root), expected)

    def test_formatting_and_parentheses_do_not_change_identity(self):
        pairs = (
            ("a==b", "a == b"),
            ("when(a==1,b==2)", "when( a == 1 , b == 2 )"),
            ('unique_by(["a","b"])', 'unique_by( [ "a" , "b" ] )'),
            ("a and b or c", "(a and b) or c"),
        )
        for first, second in pairs:
            with self.subTest(first=first):
                first_ast = validation.parse_validation_expression(first)
                second_ast = validation.parse_validation_expression(second)
                self.assertEqual(
                    validation.canonical_ast_bytes(first_ast),
                    validation.canonical_ast_bytes(second_ast),
                )
                self.assertEqual(
                    validation.semantic_ast_hash(first_ast),
                    validation.semantic_ast_hash(second_ast),
                )

    def test_hash_is_domain_separated(self):
        root = validation.parse_validation_expression("a == b", validate_builtins=False)
        undomained = "sha256:" + hashlib.sha256(
            validation.canonical_ast_bytes(root)
        ).hexdigest()
        self.assertNotEqual(validation.semantic_ast_hash(root), undomained)


class TestSecurityBoundaries(unittest.TestCase):
    def _assert_structured_limit_failure(self, source, *, limits=validation.DEFAULT_LIMITS):
        try:
            validation.parse_validation_expression(
                source,
                limits=limits,
                validate_builtins=False,
            )
        except RecursionError as error:
            self.fail(f"Raw RecursionError escaped the public parser API: {error}")
        except validation.ValidationExpressionError as error:
            self.assertEqual(error.diagnostics[0].code, validation.LIMIT_EXCEEDED)
            return error.diagnostics[0]
        self.fail("Expected a structured parser limit failure.")

    def test_structured_diagnostic_categories_are_deterministic(self):
        cases = (
            ("@", validation.LEXICAL_ERROR),
            ("", validation.SYNTAX_ERROR),
            ('"unterminated', validation.LITERAL_INVALID),
            ("9223372036854775808", validation.INTEGER_RANGE),
            ("unknown_builtin()", validation.UNKNOWN_BUILTIN),
            ("count()", validation.ARITY_MISMATCH),
            ("a * b", validation.UNSUPPORTED_SYNTAX),
        )
        for source, expected in cases:
            with self.subTest(source=source):
                with self.assertRaises(validation.ValidationExpressionError) as first:
                    validation.parse_validation_expression(source)
                with self.assertRaises(validation.ValidationExpressionError) as second:
                    validation.parse_validation_expression(source)
                self.assertEqual(first.exception.diagnostics[0].code, expected)
                self.assertEqual(
                    first.exception.diagnostics[0].as_dict(),
                    second.exception.diagnostics[0].as_dict(),
                )

    def test_negative_language_surface_is_rejected(self):
        rejected = {
            "eval(value)": validation.UNKNOWN_BUILTIN,
            "exec(value)": validation.UNKNOWN_BUILTIN,
            '__import__("os")': validation.UNSUPPORTED_SYNTAX,
            "lambda x: x": validation.UNSUPPORTED_SYNTAX,
            "import os": validation.UNSUPPORTED_SYNTAX,
            "def f()": validation.UNSUPPORTED_SYNTAX,
            "'single'": validation.UNSUPPORTED_SYNTAX,
            "1.5": validation.UNSUPPORTED_SYNTAX,
            "value[0]": validation.UNSUPPORTED_SYNTAX,
            "value[0:1]": validation.UNSUPPORTED_SYNTAX,
            "[x for x in values]": validation.UNSUPPORTED_SYNTAX,
            "a * b": validation.UNSUPPORTED_SYNTAX,
            "a / b": validation.UNSUPPORTED_SYNTAX,
            "a ** b": validation.UNSUPPORTED_SYNTAX,
            "a % b": validation.UNSUPPORTED_SYNTAX,
            "a & b": validation.UNSUPPORTED_SYNTAX,
            "a\u00a0== b": validation.LEXICAL_ERROR,
            "lookup(table = value)": validation.UNSUPPORTED_SYNTAX,
            "record.__class__": validation.UNSUPPORTED_SYNTAX,
        }
        for source, expected in rejected.items():
            with self.subTest(source=source):
                self.assertEqual(error_code(source), expected)

    def test_each_parser_limit_has_a_deterministic_diagnostic(self):
        cases = (
            (
                "abcd",
                validation.ParserLimits(max_source_length=3),
            ),
            (
                "a b",
                validation.ParserLimits(max_tokens=1),
            ),
            (
                "(((a)))",
                validation.ParserLimits(max_nesting_depth=2),
            ),
            (
                '"abc"',
                validation.ParserLimits(max_string_length=2),
            ),
            (
                "[a,b,c]",
                validation.ParserLimits(max_collection_items=2),
            ),
            (
                "a;b;c",
                validation.ParserLimits(max_sequence_statements=2),
            ),
        )
        for source, limits in cases:
            with self.subTest(source=source):
                self.assertEqual(
                    error_code(source, limits=limits), validation.LIMIT_EXCEEDED
                )

    def test_configured_nesting_limit_is_accepted(self):
        limit = validation.DEFAULT_LIMITS.max_nesting_depth
        self.assertLessEqual(limit, 64)
        parentheses = "(" * limit + "a" + ")" * limit
        calls = "when(true," * limit + "a" + ")" * limit

        self.assertIsInstance(
            validation.parse_validation_expression(
                parentheses,
                validate_builtins=False,
            ),
            validation.Identifier,
        )
        self.assertIsInstance(
            validation.parse_validation_expression(calls),
            validation.Call,
        )

    def test_configured_nesting_limit_plus_one_is_rejected_early(self):
        limit = validation.DEFAULT_LIMITS.max_nesting_depth
        sources = (
            "(" * (limit + 1) + "a" + ")" * (limit + 1),
            "when(true," * (limit + 1) + "a" + ")" * (limit + 1),
        )
        for source in sources:
            with self.subTest(prefix=source[:10]):
                diagnostic = self._assert_structured_limit_failure(source)
                context = diagnostic.as_dict()["context"]
                self.assertEqual(context["actual"], limit + 1)
                self.assertEqual(context["limit"], limit)
                self.assertEqual(context["limit_name"], "max_nesting_depth")

    def test_far_over_limit_parentheses_have_structured_failure(self):
        limit = validation.DEFAULT_LIMITS.max_nesting_depth
        depth = limit * 8
        source = "(" * depth + "a" + ")" * depth
        limits = validation.ParserLimits(
            max_tokens=len(source) + 1,
            max_nesting_depth=limit,
        )
        diagnostic = self._assert_structured_limit_failure(source, limits=limits)
        self.assertEqual(
            diagnostic.as_dict()["context"]["limit_name"],
            "max_nesting_depth",
        )

    def test_unary_stress_is_bounded_and_contains_runtime_recursion(self):
        normal_source = "not " * validation.DEFAULT_LIMITS.max_nesting_depth + "a"
        root = validation.parse_validation_expression(
            normal_source,
            validate_builtins=False,
        )
        self.assertIsInstance(root, validation.UnaryOperation)

        unary_count = sys.getrecursionlimit() * 2
        adversarial_source = "not " * unary_count + "a"
        elevated_limits = validation.ParserLimits(
            max_source_length=len(adversarial_source) + 1,
            max_tokens=unary_count + 2,
        )
        first = self._assert_structured_limit_failure(
            adversarial_source,
            limits=elevated_limits,
        )
        second = self._assert_structured_limit_failure(
            adversarial_source,
            limits=elevated_limits,
        )
        self.assertEqual(first.as_dict(), second.as_dict())
        self.assertEqual(
            first.as_dict()["context"]["limit_name"],
            "python_recursion_boundary",
        )

    def test_mixed_parentheses_calls_and_unary_stress_is_bounded(self):
        def mixed_expression(layers):
            expression = "a"
            for _ in range(layers):
                expression = f"not (when(true,{expression}))"
            return expression

        safe_layers = validation.DEFAULT_LIMITS.max_nesting_depth // 2
        safe = validation.parse_validation_expression(mixed_expression(safe_layers))
        self.assertIsInstance(safe, validation.UnaryOperation)

        diagnostic = self._assert_structured_limit_failure(
            mixed_expression(safe_layers + 1)
        )
        self.assertEqual(
            diagnostic.as_dict()["context"]["limit_name"],
            "max_nesting_depth",
        )


class TestWorkbookCorpusGolden(unittest.TestCase):
    def test_complete_active_expression_corpus(self):
        active_count, structured_only_count, corpus = active_validation_corpus()
        self.assertEqual(active_count, 308)
        self.assertEqual(len(corpus), 249)
        self.assertEqual(structured_only_count, 59)

        called_names = set()
        hashes = []
        for workbook_name, row_number, rule_id, expression in corpus:
            with self.subTest(workbook=workbook_name, row=row_number, rule=rule_id):
                first = validation.parse_validation_expression(expression)
                second = validation.parse_validation_expression(expression)
                first_hash = validation.semantic_ast_hash(first)
                self.assertRegex(first_hash, r"\Asha256:[0-9a-f]{64}\Z")
                self.assertEqual(first_hash, validation.semantic_ast_hash(second))
                hashes.append(first_hash)
                called_names.update(validation.called_builtin_names(first))

        self.assertEqual(len(hashes), 249)
        self.assertEqual(len(called_names), 54)
        self.assertEqual(called_names, validation.BUILTIN_CATALOG.names)
        self.assertEqual(called_names - validation.BUILTIN_CATALOG.names, set())


if __name__ == "__main__":
    unittest.main()
